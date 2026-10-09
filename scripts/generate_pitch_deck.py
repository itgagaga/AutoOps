# -*- coding: utf-8 -*-
"""Generate 智维超级员工 product pitch deck PPTX."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree
from pathlib import Path

TEAL_DARK = RGBColor(0x0B, 0x4F, 0x5C)
TEAL = RGBColor(0x0B, 0x72, 0x85)
TEAL_BRIGHT = RGBColor(0x14, 0x93, 0xA5)
TEAL_SOFT = RGBColor(0xE6, 0xF4, 0xF6)
SLATE = RGBColor(0x1F, 0x2A, 0x37)
SLATE_MID = RGBColor(0x52, 0x60, 0x6D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF5, 0xF8, 0xFA)
ACCENT_ORANGE = RGBColor(0xD9, 0x6A, 0x1A)
CARD_BORDER = RGBColor(0xC9, 0xDC, 0xE3)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W, H = prs.slide_width, prs.slide_height


def set_run(run, size=18, bold=False, color=SLATE, font='Microsoft YaHei'):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    r_pr = run._r.get_or_add_rPr()
    ea = r_pr.find(qn('a:ea'))
    if ea is None:
        ea = etree.SubElement(r_pr, qn('a:ea'))
    ea.set('typeface', font)


def add_rect(slide, l, t, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    return shape


def add_round(slide, l, t, w, h, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    return shape


def add_text(slide, l, t, w, h, text, size=18, bold=False, color=SLATE,
             align=PP_ALIGN.LEFT, font='Microsoft YaHei'):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    lines = text.split('\n')
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def add_bullets(slide, l, t, w, h, items, size=16, color=SLATE, spacing=8):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = '•  ' + item
        set_run(run, size=size, color=color)
    return box


def page_chrome(slide, title, subtitle=None, section=None):
    add_rect(slide, 0, 0, W, Inches(0.08), TEAL_BRIGHT)
    add_rect(slide, 0, 0, Inches(0.12), H, TEAL)
    brand = '智维超级员工  ·  Smart Ops Super Employee'
    if section:
        brand = f'{brand}  ·  {section}'
    add_text(slide, Inches(0.55), Inches(0.22), Inches(12), Inches(0.35),
             brand, size=12, color=TEAL, bold=True)
    add_text(slide, Inches(0.55), Inches(0.55), Inches(12), Inches(0.55),
             title, size=28, bold=True, color=TEAL_DARK)
    if subtitle:
        add_text(slide, Inches(0.55), Inches(1.08), Inches(12), Inches(0.4),
                 subtitle, size=15, color=SLATE_MID)
    add_rect(slide, 0, H - Inches(0.38), W, Inches(0.38), LIGHT_BG)
    add_text(slide, Inches(0.55), H - Inches(0.32), Inches(8), Inches(0.28),
             '机密 · 仅供路演演示', size=10, color=SLATE_MID)
    add_text(
        slide, Inches(10.5), H - Inches(0.32), Inches(2.5), Inches(0.28),
        'SERVER OPERATIONS', size=10, color=TEAL, align=PP_ALIGN.RIGHT, bold=True,
    )


def blank():
    return prs.slides.add_slide(prs.slide_layouts[6])


def section_divider(num, title, points):
    """Chapter opener: number + title + what this section will cover."""
    s = blank()
    add_rect(s, 0, 0, W, H, TEAL_DARK)
    add_rect(s, 0, 0, Inches(0.18), H, TEAL_BRIGHT)
    add_text(s, Inches(1.0), Inches(1.8), Inches(11), Inches(0.6),
             num, size=40, bold=True, color=TEAL_BRIGHT)
    add_text(s, Inches(1.0), Inches(2.55), Inches(11), Inches(0.7),
             title, size=36, bold=True, color=WHITE)
    add_text(s, Inches(1.0), Inches(3.4), Inches(11), Inches(0.4),
             '本章将讲清楚', size=14, color=RGBColor(0xA8, 0xD4, 0xDC))
    add_bullets(s, Inches(1.0), Inches(3.95), Inches(11), Inches(2.2),
                points, size=16, color=WHITE, spacing=10)
    return s


def two_col_cards(slide, cards, top=1.65):
    for i, (t, items) in enumerate(cards):
        x = Inches(0.55 + (i % 2) * 6.3)
        y = Inches(top + (i // 2) * 2.45)
        add_round(slide, x, y, Inches(6.0), Inches(2.25), WHITE, CARD_BORDER)
        add_rect(slide, x, y, Inches(0.12), Inches(2.25), TEAL_BRIGHT if i % 2 == 0 else ACCENT_ORANGE)
        add_text(slide, x + Inches(0.35), y + Inches(0.25), Inches(5.4), Inches(0.4),
                 t, size=18, bold=True, color=TEAL_DARK)
        add_bullets(slide, x + Inches(0.35), y + Inches(0.75), Inches(5.4), Inches(1.3),
                    items, size=13, color=SLATE_MID, spacing=4)


# ========== Cover ==========
s = blank()
add_rect(s, 0, 0, W, H, TEAL_DARK)
add_rect(s, 0, Inches(5.55), W, Inches(1.95), TEAL)
add_text(s, Inches(0.9), Inches(1.6), Inches(11), Inches(0.4),
         'SERVER OPERATIONS', size=16, bold=True, color=TEAL_BRIGHT)
add_text(s, Inches(0.9), Inches(2.15), Inches(11.5), Inches(0.9),
         '智维超级员工', size=48, bold=True, color=WHITE)
add_text(s, Inches(0.9), Inches(3.15), Inches(11.5), Inches(0.7),
         '面向服务器运维的多 Agent 全流程自治平台', size=22, color=RGBColor(0xC8, 0xE6, 0xEC))
add_text(s, Inches(0.9), Inches(4.1), Inches(11.5), Inches(0.5),
         '从日志异常到可部署制品 —— 带审批门禁的真实修码闭环', size=16, color=RGBColor(0xA8, 0xD4, 0xDC))
add_text(s, Inches(0.9), Inches(5.85), Inches(8), Inches(0.4),
         '产品路演  ·  Pitch Deck', size=18, bold=True, color=WHITE)
add_text(s, Inches(0.9), Inches(6.35), Inches(10), Inches(0.35),
         'AgentScope-Java  ·  Vue3  ·  Cursor / CodeX / Claude CLI', size=13,
         color=RGBColor(0xD0, 0xEE, 0xF2))

# ========== Agenda ==========
s = blank()
page_chrome(s, '目录', '路演议程 · 每章含章节页与详细介绍')
agenda = [
    ('01', '为什么现在做', '行业窗口、运维痛点与时机'),
    ('02', '产品定位与价值主张', '超级员工定位与四大能力主张'),
    ('03', '核心能力与架构', '四层架构、功能全景与技术栈'),
    ('04', '多 Agent 协作与闭环', '角色分工与端到端自治流程'),
    ('05', '安全可控与差异化', '治理门禁与竞品差异对比'),
    ('06', '演示场景与落地路径', '现场演示剧本与推广节奏'),
]
for i, (num, title, desc) in enumerate(agenda):
    col, row = i % 3, i // 3
    x = Inches(0.55 + col * 4.2)
    y = Inches(1.65 + row * 2.4)
    add_round(s, x, y, Inches(4.0), Inches(2.15), WHITE, CARD_BORDER)
    add_text(s, x + Inches(0.25), y + Inches(0.25), Inches(3.5), Inches(0.45),
             num, size=26, bold=True, color=TEAL_BRIGHT)
    add_text(s, x + Inches(0.25), y + Inches(0.85), Inches(3.5), Inches(0.45),
             title, size=16, bold=True, color=SLATE)
    add_text(s, x + Inches(0.25), y + Inches(1.4), Inches(3.5), Inches(0.5),
             desc, size=12, color=SLATE_MID)

# ========== 01 为什么现在做 ==========
section_divider('01', '为什么现在做', [
    '服务器运维故障处置为什么难规模化',
    '传统平台与 ChatOps 各自卡在哪里',
    '为什么「多 Agent + 真实修码 + 审批门禁」此刻可落地',
])

s = blank()
page_chrome(s, '市场与时机', '运维自动化进入「可改码」阶段', section='01 为什么现在做')
two_col_cards(s, [
    ('窗口已经打开', [
        '大模型与 Coding Agent 已能产出可审阅的真实代码变更',
        '企业需要的不只是告警，而是缩短 MTTR 的闭环能力',
        '专家经验稀缺，组织需要把处置路径产品化、可复制',
    ]),
    ('旧方案到了瓶颈', [
        '监控/工单平台：发现快，修复仍靠人',
        '脚本自动化：覆盖固定场景，难应对新异常',
        '纯 ChatOps：建议多、约束少，生产不敢放权',
    ]),
    ('新范式的关键三件套', [
        '多 Agent 分工：证据、根因、修复、验证各司其职',
        'CLI 修码引擎：Cursor / CodeX / Claude 产出真实 Diff',
        '状态机门禁：审批、指纹绑定、审计，保证可控',
    ]),
    ('适合优先切入的场景', [
        '多技术栈应用运行时异常（日志可观测）',
        '有代码工作区、可本地验证构建的研发运维环境',
        '需要「人在回路」审批，而不是全自动裸奔上线',
    ]),
])

s = blank()
page_chrome(s, '行业痛点详解', '四类结构性矛盾，决定了必须换范式', section='01 为什么现在做')
pains = [
    ('痛点一 · 证据分散', '故障处置高度依赖专家翻日志、猜根因；证据不成链，复盘难沉淀，MTTR 居高不下。'),
    ('痛点二 · 修码断层', '平台多停留在告警与工单；从「知道坏了」到「改出补丁」仍靠人工，链路长、易漏。'),
    ('痛点三 · 建议难落地', 'ChatOps 常只有对话建议：无状态机、无补丁指纹、无强制审批，副作用风险高。'),
    ('痛点四 · 安全不敢放', '若 Agent 可任意 Shell，生产不可接受；需要只读诊断 + 工作区沙箱 + 双段审批。'),
]
for i, (t, d) in enumerate(pains):
    x = Inches(0.55 + (i % 2) * 6.3)
    y = Inches(1.65 + (i // 2) * 2.4)
    add_round(s, x, y, Inches(6.0), Inches(2.2), WHITE, CARD_BORDER)
    add_rect(s, x, y, Inches(0.12), Inches(2.2), TEAL_BRIGHT if i % 2 == 0 else ACCENT_ORANGE)
    add_text(s, x + Inches(0.4), y + Inches(0.35), Inches(5.3), Inches(0.45),
             t, size=18, bold=True, color=TEAL_DARK)
    add_text(s, x + Inches(0.4), y + Inches(0.95), Inches(5.3), Inches(0.95),
             d, size=14, color=SLATE_MID)

# ========== 02 产品定位与价值主张 ==========
section_divider('02', '产品定位与价值主张', [
    '一句话讲清「智维超级员工」是什么、不是什么',
    '四大价值主张：看得清 / 想得透 / 改得动 / 控得住',
    '对客户组织带来的可量化业务价值',
])

s = blank()
page_chrome(s, '产品是什么', '可审计的运维超级员工，而非聊天机器人', section='02 产品定位与价值主张')
add_round(s, Inches(0.55), Inches(1.65), Inches(12.2), Inches(1.55), TEAL_SOFT)
add_text(
    s, Inches(0.9), Inches(1.9), Inches(11.5), Inches(1.1),
    '「智维超级员工」面向服务器运维故障自愈，把\n'
    '「拉日志 → 分析 → 本地修码 → 验证 → 审批 → 构建/部署」串成可控闭环。',
    size=18, bold=True, color=TEAL_DARK,
)
compare = [
    ('它是', [
        '带状态机约束的多 Agent 运维平台',
        '能调用 Coding CLI 产出真实 Diff',
        '人在回路审批后的受控自治系统',
    ]),
    ('它不是', [
        '只会对话给建议的 ChatOps 机器人',
        '放开任意 Shell 的无人值守脚本机',
        '只做监控告警、不碰代码的工单系统',
    ]),
]
for i, (t, items) in enumerate(compare):
    x = Inches(0.55 + i * 6.3)
    add_round(s, x, Inches(3.5), Inches(6.0), Inches(2.85), WHITE, CARD_BORDER)
    add_rect(s, x, Inches(3.5), Inches(6.0), Inches(0.65), TEAL if i == 0 else SLATE_MID)
    add_text(s, x + Inches(0.3), Inches(3.62), Inches(5.4), Inches(0.45),
             t, size=18, bold=True, color=WHITE)
    add_bullets(s, x + Inches(0.35), Inches(4.4), Inches(5.4), Inches(1.7),
                items, size=15, color=SLATE, spacing=8)

s = blank()
page_chrome(s, '四大价值主张', '从「看见问题」到「安全交卷」一条链打通', section='02 产品定位与价值主张')
claims = [
    ('看得清', '巡检发现', 'SSH 只读模板拉日志；错误阈值触发故障任务；证据可追溯。'),
    ('想得透', '智能研判', '协调 Agent + 子智能体分工；根因分析可结合知识库线索。'),
    ('改得动', '真实修码', 'Cursor / CodeX / Claude CLI 在工作区改码；产出可审阅 Diff。'),
    ('控得住', '治理门禁', '补丁指纹绑定验证/部署；双段人工审批；工具调用全程审计。'),
]
for i, (t, tag, d) in enumerate(claims):
    x = Inches(0.45 + i * 3.2)
    add_round(s, x, Inches(1.7), Inches(3.05), Inches(4.9), WHITE, CARD_BORDER)
    add_rect(s, x, Inches(1.7), Inches(3.05), Inches(1.35), TEAL_DARK if i % 2 == 0 else TEAL)
    add_text(s, x + Inches(0.15), Inches(1.9), Inches(2.75), Inches(0.45),
             t, size=22, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.15), Inches(2.4), Inches(2.75), Inches(0.4),
             tag, size=13, color=TEAL_SOFT, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.25), Inches(3.4), Inches(2.55), Inches(2.6),
             d, size=14, color=SLATE)

s = blank()
page_chrome(s, '客户价值', '把专家经验产品化，把处置链路工业化', section='02 产品定位与价值主张')
vals = [
    ('缩短 MTTR', [
        '异常自动发现与证据收集，减少「找日志」时间',
        '根因到补丁路径标准化，降低临场发挥成本',
    ]),
    ('降低专家依赖', [
        '知识库沉淀历史案例，新人可按门禁闭环处置',
        '多 Agent 分工把复杂任务拆成可执行步骤',
    ]),
    ('提升变更安全', [
        '真实 Diff 可审阅，避免「黑盒改生产」',
        '验证绑定补丁指纹，减少错包、错部署',
    ]),
    ('可规模复制', [
        '主机 / 凭据 / 运维任务模板化配置',
        '三引擎可切换，适配达梦等国产数据库',
    ]),
]
two_col_cards(s, vals)

# ========== 03 核心能力与架构 ==========
section_divider('03', '核心能力与架构', [
    '平台四层架构如何支撑自治闭环',
    '面向运维人员的功能全景有哪些入口',
    '关键技术组件：运行时、修复网关、制品构建',
])

s = blank()
page_chrome(s, '整体架构', '体验 · 编排 · 能力 · 资源 四层协同', section='03 核心能力与架构')
layers = [
    ('体验层', '运维工作台 / Agent 中心 / 资源中心 / SSH·DB·Git 客户端 / 知识库', TEAL_BRIGHT),
    ('编排层', '协调 Agent + 子智能体 · Checkpoint / 租约 / 心跳 / 审批后恢复执行', TEAL),
    ('能力层', '日志巡检 · 生命周期状态机 · 三引擎 CLI 修复 · 多栈验证 · 混合部署', TEAL_DARK),
    ('资源层', '主机指纹与日志白名单 · 凭据保险库 · 达梦 / MySQL / Oracle / Kingbase', SLATE),
]
for i, (name, desc, color) in enumerate(layers):
    y = Inches(1.7 + i * 1.2)
    add_round(s, Inches(0.55), y, Inches(12.2), Inches(1.05), WHITE, CARD_BORDER)
    add_round(s, Inches(0.75), y + Inches(0.22), Inches(1.8), Inches(0.6), color)
    add_text(s, Inches(0.75), y + Inches(0.32), Inches(1.8), Inches(0.45),
             name, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, Inches(2.8), y + Inches(0.32), Inches(9.5), Inches(0.5),
             desc, size=16, color=SLATE)

s = blank()
page_chrome(s, '功能全景', '日常运维入口与 Agent 自治能力一体呈现', section='03 核心能力与架构')
mods = [
    ('运维工作台', [
        '运维任务：主机、日志模板、巡检周期',
        '工程画像：多栈后端 + 端侧类型',
        '工作区、修复引擎与故障启停',
    ]),
    ('Agent 中心', [
        'Agent 管理与三引擎环境检测',
        '工具管理与授权策略配置',
        '执行记录追溯与审批恢复',
    ]),
    ('资源中心', [
        '主机纳管：指纹 + 日志白名单',
        '凭据保险库（密钥不回显）',
        '数据库等资源统一登记',
    ]),
    ('知识库', [
        '文档导入 + Agent 预生成草稿',
        '历史案例作为研判线索',
        '不可用时降级，不阻断主流程',
    ]),
    ('运维客户端', [
        'SSH 终端：人工排障协同',
        '数据库终端 / Git 客户端',
        '人机协同，而非替代人工',
    ]),
    ('验证与部署', [
        '多栈验证与制品构建',
        '路径替换 / WAR / Docker / Web 静态',
        '安全下载与受控部署确认',
    ]),
]
for i, (t, items) in enumerate(mods):
    col, row = i % 3, i // 3
    x, y = Inches(0.45 + col * 4.25), Inches(1.6 + row * 2.55)
    add_round(s, x, y, Inches(4.05), Inches(2.35), WHITE, CARD_BORDER)
    add_text(s, x + Inches(0.25), y + Inches(0.25), Inches(3.55), Inches(0.4),
             t, size=16, bold=True, color=TEAL)
    add_bullets(s, x + Inches(0.25), y + Inches(0.75), Inches(3.55), Inches(1.4),
                items, size=12, color=SLATE_MID, spacing=4)

s = blank()
page_chrome(s, '关键技术能力', '支撑「改得动、控得住」的三块底座', section='03 核心能力与架构')
techs = [
    ('Agent 运行时', [
        '基于 AgentScope ReAct 多智能体编排',
        'Checkpoint / 租约 / 心跳 / 超时回收',
        '失败分类重试；审批通过后恢复执行',
        '工具调用审计，权限按策略收敛',
    ]),
    ('CLI 修复网关', [
        '统一抽象：CodeX / Cursor / Claude',
        '可固定引擎，也可 auto 失败切换',
        '工作区边界校验，越界即失败',
        'git diff 留痕，补丁可审阅可回放',
    ]),
    ('验证与部署', [
        '按工程画像验证，绑定真实补丁差异',
        'Java / Node / Go / Python / Docker 多栈',
        '混合部署：后端替换 + Web 静态同步',
        '部署前再次审批，人工确认完结',
    ]),
]
for i, (t, items) in enumerate(techs):
    x = Inches(0.5 + i * 4.2)
    add_round(s, x, Inches(1.7), Inches(4.0), Inches(4.9), WHITE, CARD_BORDER)
    add_rect(s, x, Inches(1.7), Inches(4.0), Inches(0.85), TEAL_DARK)
    add_text(s, x + Inches(0.2), Inches(1.9), Inches(3.6), Inches(0.5),
             t, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_bullets(s, x + Inches(0.3), Inches(2.9), Inches(3.5), Inches(3.3),
                items, size=14, color=SLATE, spacing=10)

# ========== 04 多 Agent 协作与闭环 ==========
section_divider('04', '多 Agent 协作与闭环', [
    '各 Agent 角色如何分工协作',
    '从巡检到制品的端到端状态流转',
    '门禁点在哪里、人在回路如何介入',
])

s = blank()
page_chrome(s, '多 Agent 角色分工', '证据驱动，避免「一个大模型包打天下」', section='04 多 Agent 协作与闭环')
agents = [
    ('协调 Agent', '总控', '任务拆解与调度\n门禁判断与恢复\n串联子智能体'),
    ('证据收集', '现场', 'SSH 只读诊断\n拉取日志与证据\n形成可追溯材料'),
    ('根因分析', '研判', '关联知识库线索\n锁定可疑根因\n输出分析结论'),
    ('修复规划', '方案', '生成补丁计划\n绑定补丁指纹\n进入预览门禁'),
    ('验证 / 部署', '交卷', '多栈验证\n混合部署\n受控确认'),
]
for i, (t, tag, d) in enumerate(agents):
    x = Inches(0.35 + i * 2.55)
    add_round(s, x, Inches(1.7), Inches(2.45), Inches(4.9), WHITE, CARD_BORDER)
    add_rect(s, x, Inches(1.7), Inches(2.45), Inches(1.15), TEAL if i else TEAL_DARK)
    add_text(s, x + Inches(0.1), Inches(1.85), Inches(2.25), Inches(0.4),
             t, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.1), Inches(2.3), Inches(2.25), Inches(0.35),
             tag, size=12, color=TEAL_SOFT, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.15), Inches(3.2), Inches(2.15), Inches(2.8),
             d, size=13, color=SLATE, align=PP_ALIGN.CENTER)

s = blank()
page_chrome(s, '端到端自治闭环', '七步串成可审计、可中断、可恢复的流水线', section='04 多 Agent 协作与闭环')
steps = [
    ('1', '启用巡检', '运维任务绑定主机\n与代码工作区'),
    ('2', '拉日志', 'SSH 只读模板\n错误阈值判异常'),
    ('3', '多 Agent', '证据 / 根因\n修复规划协作'),
    ('4', '补丁预览', '计划 + 指纹\n进入审批门禁'),
    ('5', 'CLI 修码', '真实 Diff\n工作区沙箱执行'),
    ('6', '验证构建', '按画像验证\n产出可部署制品'),
    ('7', '部署完结', '再次审批\n混合部署/沉淀'),
]
for i, (n, t, d) in enumerate(steps):
    x = Inches(0.3 + i * 1.85)
    add_round(s, x, Inches(1.75), Inches(1.75), Inches(3.5), WHITE, CARD_BORDER)
    circle = s.shapes.add_shape(
        MSO_SHAPE.OVAL, x + Inches(0.6), Inches(2.0), Inches(0.5), Inches(0.5),
    )
    circle.fill.solid()
    circle.fill.fore_color.rgb = TEAL_BRIGHT
    circle.line.fill.background()
    add_text(s, x + Inches(0.6), Inches(2.05), Inches(0.5), Inches(0.4),
             n, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.1), Inches(2.7), Inches(1.55), Inches(0.45),
             t, size=13, bold=True, color=TEAL_DARK, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.1), Inches(3.3), Inches(1.55), Inches(1.5),
             d, size=11, color=SLATE_MID, align=PP_ALIGN.CENTER)

add_round(s, Inches(0.55), Inches(5.5), Inches(12.2), Inches(1.15), TEAL_SOFT)
add_text(
    s, Inches(0.85), Inches(5.7), Inches(11.6), Inches(0.8),
    '人在回路关键点：① 本地修码前审批　② 部署执行前再审批\n'
    '系统约束：验证/部署必须绑定真实补丁指纹；Checkpoint 支持中断后恢复。',
    size=14, bold=True, color=TEAL_DARK,
)

# ========== 05 安全可控与差异化 ==========
section_divider('05', '安全可控与差异化', [
    '安全边界如何设计：远程只读、本地沙箱、双段审批',
    '相对传统运维平台与纯 ChatOps 的差异矩阵',
    '为什么企业敢把「修码」交给 Agent',
])

s = blank()
page_chrome(s, '安全与治理详解', '让 Agent「能干活，但不能乱干活」', section='05 安全可控与差异化')
secs = [
    ('远程只读', [
        '固定 SSH 诊断模板，禁止任意命令',
        '主机密钥指纹校验，防中间人',
        '应用日志路径白名单约束',
    ]),
    ('本地沙箱', [
        '修码严格限定代码工作区',
        'Prompt 禁 commit / push / 裸部署',
        'git diff 作为变更唯一审计源',
    ]),
    ('双段审批', [
        '本地修复前必须人工批准',
        '部署执行前再次人工确认',
        '工具调用与审批全程留痕',
    ]),
    ('引擎可控', [
        'Cursor / CodeX / Claude 可切换',
        '支持 auto 失败切换与降级',
        '授权策略按角色收敛工具面',
    ]),
]
for i, (t, items) in enumerate(secs):
    x = Inches(0.45 + i * 3.2)
    add_round(s, x, Inches(1.7), Inches(3.05), Inches(4.9), WHITE, CARD_BORDER)
    add_rect(s, x, Inches(1.7), Inches(3.05), Inches(0.85), TEAL_DARK)
    add_text(s, x + Inches(0.15), Inches(1.9), Inches(2.75), Inches(0.5),
             t, size=17, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_bullets(s, x + Inches(0.25), Inches(2.9), Inches(2.6), Inches(3.2),
                items, size=13, color=SLATE, spacing=10)

s = blank()
page_chrome(s, '差异化对比', '相对传统运维平台与纯 ChatOps', section='05 安全可控与差异化')
headers = ['维度', '传统运维平台', '纯 ChatOps', '智维超级员工']
rows = [
    ['处置深度', '告警 + 工单', '对话建议', '修码到验证到制品'],
    ['过程约束', '流程靠制度', '弱约束', '状态机 + 指纹门禁'],
    ['执行形态', '脚本 / 人工', '建议人工执行', '三引擎 CLI 真实改码'],
    ['安全边界', '权限模型不一', '易越权', '只读模板 + 双段审批'],
    ['可恢复性', '依赖人工续作', '会话易断', 'Checkpoint / 租约恢复'],
]
for i, h in enumerate(headers):
    x = Inches(0.45 + i * 3.15)
    add_rect(s, x, Inches(1.65), Inches(3.05), Inches(0.55), TEAL_DARK if i == 3 else TEAL)
    add_text(s, x, Inches(1.75), Inches(3.05), Inches(0.4),
             h, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
for r, row in enumerate(rows):
    y = Inches(2.25 + r * 0.78)
    for c, cell in enumerate(row):
        x = Inches(0.45 + c * 3.15)
        bg = TEAL_SOFT if c == 3 else (WHITE if r % 2 == 0 else LIGHT_BG)
        add_rect(s, x, y, Inches(3.05), Inches(0.72), bg, CARD_BORDER)
        add_text(
            s, x + Inches(0.1), y + Inches(0.18), Inches(2.85), Inches(0.4), cell,
            size=13, bold=(c == 0 or c == 3),
            color=TEAL_DARK if c == 3 else SLATE, align=PP_ALIGN.CENTER,
        )

# ========== 06 演示场景与落地路径 ==========
section_divider('06', '演示场景与落地路径', [
    '用测试工程做端到端自愈演示的剧本',
    '可展示的缺陷类型与验收关注点',
    '从样板闭环到规模推广的三阶段路径',
])

s = blank()
page_chrome(s, '演示场景详解', 'demo-ops-app：故意埋缺陷，验证自治闭环', section='06 演示场景与落地路径')
add_round(s, Inches(0.55), Inches(1.65), Inches(7.7), Inches(5.0), WHITE, CARD_BORDER)
add_text(s, Inches(0.9), Inches(1.95), Inches(7.1), Inches(0.4),
         '现场演示剧本', size=18, bold=True, color=TEAL_DARK)
add_bullets(s, Inches(0.9), Inches(2.55), Inches(7.1), Inches(3.8), [
    '启动演示应用，FaultInjector 定时制造 ERROR 日志',
    '在运维工作台创建任务：绑定目标主机与代码工作区',
    '选择修复引擎（Cursor / CodeX / Claude 或 auto）',
    '启用巡检 → 自动发现异常并派生故障事件',
    '观察多 Agent：证据收集 → 根因 → 补丁预览',
    '审批通过后 CLI 本地修码，审阅真实 Diff',
    '按工程画像验证并完成受控部署，确认完结'
], size=14, spacing=8)

add_round(s, Inches(8.5), Inches(1.65), Inches(4.25), Inches(5.0), TEAL_DARK)
add_text(s, Inches(8.8), Inches(2.0), Inches(3.7), Inches(0.4),
         '可演示缺陷类型', size=16, bold=True, color=TEAL_BRIGHT)
add_bullets(s, Inches(8.8), Inches(2.65), Inches(3.7), Inches(3.5), [
    '空指针 NPE',
    '数字格式异常',
    '日期解析失败',
    '分页越界',
    '文件不存在',
    '字符串越界',
    '并发修改异常',
], size=14, color=WHITE, spacing=8)

s = blank()
page_chrome(s, '演示验收关注点', '路演时建议重点展示的「可信证据」', section='06 演示场景与落地路径')
focus = [
    ('过程可见', [
        '执行记录与时间线完整可回放',
        'Agent 每步输出与工具调用可审计',
        '故障状态机阶段清晰可解释',
    ]),
    ('变更可信', [
        '补丁预览与真实 Diff 一致',
        '验证绑定补丁指纹，防偷换',
        '构建产物可下载、可对照版本',
    ]),
    ('风险可控', [
        '远程侧仅只读诊断模板',
        '修码前后均有人工审批节点',
        '拒绝越权路径与越界改文件',
    ]),
]
for i, (t, items) in enumerate(focus):
    x = Inches(0.5 + i * 4.2)
    add_round(s, x, Inches(1.7), Inches(4.0), Inches(4.9), WHITE, CARD_BORDER)
    add_rect(s, x, Inches(1.7), Inches(4.0), Inches(0.85), TEAL)
    add_text(s, x + Inches(0.2), Inches(1.9), Inches(3.6), Inches(0.5),
             t, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_bullets(s, x + Inches(0.3), Inches(2.9), Inches(3.5), Inches(3.3),
                items, size=14, color=SLATE, spacing=12)

s = blank()
page_chrome(s, '落地路径', '先打通一条 Java 应用自愈样板，再横向扩展', section='06 演示场景与落地路径')
phases = [
    ('Phase 1', '样板闭环', [
        '选定 1–2 个可观测日志的应用（建议含 Java 或 Web）',
        '完成主机、凭据、白名单配置',
        '打通巡检 → 修码 → 验证 → 审批',
        '形成可复现演示与操作手册',
    ]),
    ('Phase 2', '规则与知识', [
        '沉淀高频故障案例到知识库',
        '细化授权策略与工具白名单',
        '统一修复引擎与审批角色',
        '建立复盘与效果度量指标',
    ]),
    ('Phase 3', '规模推广', [
        '多主机、多任务并行巡检',
        '扩展更多应用与日志模板',
        '对接现有部署/发布体系',
        '形成组织级运维自治能力',
    ]),
]
for i, (p, t, items) in enumerate(phases):
    x = Inches(0.5 + i * 4.2)
    add_round(s, x, Inches(1.7), Inches(4.0), Inches(4.9), WHITE, CARD_BORDER)
    add_rect(s, x, Inches(1.7), Inches(4.0), Inches(1.15), TEAL if i < 2 else TEAL_DARK)
    add_text(s, x + Inches(0.2), Inches(1.85), Inches(3.6), Inches(0.35),
             p, size=12, bold=True, color=TEAL_SOFT, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.2), Inches(2.25), Inches(3.6), Inches(0.4),
             t, size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_bullets(s, x + Inches(0.3), Inches(3.2), Inches(3.5), Inches(3.0),
                items, size=13, color=SLATE, spacing=10)

# ========== Closing ==========
s = blank()
add_rect(s, 0, 0, W, H, TEAL_DARK)
add_rect(s, 0, Inches(5.8), W, Inches(1.7), TEAL)
add_text(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(0.7),
         '让每一次故障处置，都变成可复用的自治能力', size=28, bold=True, color=WHITE)
add_text(
    s, Inches(0.9), Inches(2.9), Inches(11.5), Inches(0.9),
    '智维超级员工 —— 带审批门禁与真实修码引擎的服务器运维超级员工',
    size=18, color=RGBColor(0xC8, 0xE6, 0xEC),
)
add_text(s, Inches(0.9), Inches(4.1), Inches(11.5), Inches(0.5),
         '谢谢 · 欢迎交流与现场演示', size=20, bold=True, color=TEAL_BRIGHT)
add_text(s, Inches(0.9), Inches(6.25), Inches(11.5), Inches(0.5),
         'Q & A', size=28, bold=True, color=WHITE)

root = Path(__file__).resolve().parents[1]
candidates = [
    root / '智维超级员工-产品路演.pptx',
    root / '智维超级员工-产品路演-v2.pptx',
]
out = None
last_err = None
for path in candidates:
    try:
        prs.save(str(path))
        out = path
        break
    except PermissionError as exc:
        last_err = exc
if out is None:
    raise last_err
print(out)
print('slides:', len(prs.slides))
