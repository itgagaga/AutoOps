# -*- coding: utf-8 -*-
"""Generate 智维超级员工操作手册（运维/管理人员）."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TEAL = RGBColor(0x0B, 0x72, 0x85)
TEAL_DARK = RGBColor(0x0B, 0x4F, 0x5C)
SLATE = RGBColor(0x1F, 0x2A, 0x37)
GRAY = RGBColor(0x52, 0x60, 0x6D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run_font(run, size=11, bold=False, color=SLATE, name='微软雅黑'):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)


def add_heading_cn(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(
            run,
            size={1: 18, 2: 14, 3: 12}.get(level, 11),
            bold=True,
            color=TEAL_DARK if level == 1 else TEAL,
        )
    return p


def add_para(doc, text, size=11, bold=False, color=SLATE, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_bullets(doc, items, size=11):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.clear()
        run = p.add_run(item)
        set_run_font(run, size=size, color=SLATE)
        p.paragraph_format.space_after = Pt(3)


def add_numbered(doc, items, size=11):
    for item in items:
        p = doc.add_paragraph(style='List Number')
        p.clear()
        run = p.add_run(item)
        set_run_font(run, size=size, color=SLATE)
        p.paragraph_format.space_after = Pt(4)


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size=9, color=SLATE, name='Consolas')
    run.font.name = 'Consolas'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), 'F5F8FA')
    shd.set(qn('w:val'), 'clear')
    p._p.get_or_add_pPr().append(shd)
    return p


def set_cell_text(cell, text, bold=False, color=SLATE, size=10):
    cell.text = ''
    p = cell.paragraphs[0]
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def shade_header_row(row):
    for cell in row.cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), '0B7285')
        shd.set(qn('w:val'), 'clear')
        tc_pr.append(shd)
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = WHITE
                run.font.bold = True


def add_table(doc, headers, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, size=10)
    shade_header_row(table.rows[0])
    for r_i, row in enumerate(rows, start=1):
        for c_i, val in enumerate(row):
            set_cell_text(table.rows[r_i].cells[c_i], val, size=10)
    doc.add_paragraph()
    return table


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    # Cover
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(36)
    r = t.add_run('智维超级员工')
    set_run_font(r, size=28, bold=True, color=TEAL_DARK)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run('运维与管理人员操作手册')
    set_run_font(r, size=20, bold=True, color=TEAL)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run('面向服务器运维的多 Agent 全流程自治平台')
    set_run_font(r, size=12, color=GRAY)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(16)
    r = meta.add_run('版本：V1.2 ｜ 适用对象：运维人员、系统/业务管理人员 ｜ 内部资料')
    set_run_font(r, size=10, color=GRAY)

    add_para(
        doc,
        '本手册说明平台日常配置、巡检运维、故障处置、审批治理与安全规范。'
        '读者无需具备 Agent 开发背景，按角色完成对应操作即可。',
        size=10,
        color=GRAY,
    )

    # 1
    add_heading_cn(doc, '一、产品概述', 1)
    add_para(
        doc,
        '智维超级员工用于服务器应用运维故障的受控自治处置：系统按运维任务周期性拉取日志，'
        '发现异常后由多 Agent 完成证据收集、根因分析与修复规划；在人工审批通过后，'
        '调用本地修复引擎（Cursor / CodeX / Claude CLI）在指定代码工作区生成真实补丁，'
        '再按工程画像完成验证、制品构建与受控部署（支持多技术栈与前后端混合部署）完成闭环。',
    )
    add_para(doc, '核心原则：', bold=True)
    add_bullets(doc, [
        '远程诊断只读：通过固定 SSH 模板取证，禁止任意 Shell。',
        '本地修码受控：限定代码工作区，变更可 Diff、可审计。',
        '人在回路：修码与部署等有副作用操作必须审批。',
        '状态机约束：阶段乱序或重复操作会被系统拒绝或幂等处理。',
    ])

    # 2
    add_heading_cn(doc, '二、角色与职责', 1)
    add_table(
        doc,
        ['角色', '典型人员', '主要职责'],
        [
            ('系统管理员', '平台管理员', '账号权限、授权策略、系统配置、全局审计'),
            ('资源管理员', '运维负责人', '主机、凭据、数据库等资源纳管与启用'),
            ('运维人员', '一线/二线运维', '运维任务配置、巡检监控、故障跟进、协同排障'),
            ('审批人员', '运维主管/研发负责人', '补丁修复审批、部署审批、方案驳回'),
            ('只读人员', '业务/审计相关方', '查看任务、执行记录与处置结果（按授权）'),
        ],
    )
    add_para(doc, '说明：实际菜单与按钮以账号权限为准；无权限时请联系系统管理员授权。', size=10, color=GRAY)

    # 3
    add_heading_cn(doc, '三、登录与界面导览', 1)
    add_heading_cn(doc, '3.1 登录', 2)
    add_numbered(doc, [
        '使用浏览器打开平台地址（由部署方提供）。',
        '输入账号、密码及验证码登录。',
        '首次登录建议立即修改初始密码；勿将管理员口令长期用于日常运维号。',
    ])

    add_heading_cn(doc, '3.2 界面结构', 2)
    add_bullets(doc, [
        '顶部：品牌区、产品说明书入口、当前用户（修改密码/退出）。',
        '顶部标签栏：已打开功能页，可切换、关闭；支持拖拽调整标签顺序。',
        '左侧菜单：运维工作台、Agent 中心、资源中心、SSH/数据库/Git 客户端、系统管理等。',
        '主工作区：当前功能页内容（列表、表单、详情、终端等）。',
    ])

    add_heading_cn(doc, '3.3 常用功能入口', 2)
    add_table(
        doc,
        ['功能', '菜单位置', '用途'],
        [
            ('运维工作台', '左侧「运维工作台」', '运维任务、巡检、故障总览'),
            ('执行记录', 'Agent 中心 / 标签页', '查看 Agent 执行过程与审计'),
            ('主机管理', '资源中心', '主机指纹、日志白名单、连通性'),
            ('凭据管理', '资源中心', 'SSH/数据库/Git 等凭据'),
            ('工具管理 / 授权策略', 'Agent 中心', '工具面与访问策略'),
            ('知识库', '知识库菜单', '案例沉淀；支持文档导入与 AI 预生成'),
            ('SSH / 数据库 / Git', '对应客户端菜单', '人工排障与协同'),
            ('Agent 管理', 'Agent 中心', 'Agent 配置、状态与 CodeX/Cursor/Claude 环境检测'),
        ],
    )

    # 4
    add_heading_cn(doc, '四、资源准备（管理人员 / 资源管理员）', 1)
    add_para(doc, '在创建运维任务前，必须完成凭据与主机配置，否则无法拉日志或启用巡检。')

    add_heading_cn(doc, '4.1 凭据管理', 2)
    add_numbered(doc, [
        '进入「凭据管理」→「新增」。',
        '选择类型（SSH / 数据库 / Git 等），填写名称、用户名及密钥或密码。',
        '保存后密钥不会在页面回显；编辑主机时从下拉选择已建凭据。',
        '若新建后下拉看不到，重新打开主机新增/编辑窗口以刷新选项。',
    ])
    add_para(doc, '管理要求：', bold=True)
    add_bullets(doc, [
        '生产凭据与测试凭据分离；最小权限原则。',
        '人员离职或轮岗时及时禁用/轮换凭据。',
        '禁止在聊天工具中明文传递密钥。',
    ])

    add_heading_cn(doc, '4.2 主机管理', 2)
    add_numbered(doc, [
        '进入「主机管理」→「新增/编辑」。',
        '填写编码、名称、IP、SSH 端口、操作系统、环境，并绑定 SSH 凭据。',
        '配置「SSH 主机密钥指纹」列表（SHA256:...）。获取示例：',
    ])
    add_code(doc, 'ssh-keyscan -p <端口> <主机IP> 2>$null | ssh-keygen -lf - -E sha256')
    add_numbered(doc, [
        '配置「允许读取的日志路径」白名单：可为日志文件或目录绝对路径。',
        '建议先保存为未启用 → 点击「探测连通性」→ 通过后再启用。',
        '启用时系统会再次探测，失败则拒绝启用。',
    ])
    add_para(doc, '日志白名单规则：', bold=True)
    add_bullets(doc, [
        '必须使用绝对路径；禁止相对路径、`..`、空格及 Shell 元字符。',
        '路径最后一段含扩展名（如 app.log）按文件处理；否则按目录处理。',
        '目录模式下系统优先选取含当日日期的日志，否则取最新匹配日志再读取。',
        '路径需符合平台安全校验；含中文路径段可能导致白名单校验失败，建议使用英文路径或 WSL 映射路径。',
    ])

    add_heading_cn(doc, '4.3 其他资源', 2)
    add_bullets(doc, [
        '数据库资源：供数据库终端及授权访问使用，按环境登记连接信息与凭据。',
        '资源启用前确认网络策略、账号权限与变更窗口要求。',
    ])

    # 5
    add_heading_cn(doc, '五、运维任务管理（运维人员）', 1)
    add_heading_cn(doc, '5.1 新建运维任务', 2)
    add_numbered(doc, [
        '进入「运维工作台」→「新增运维任务」。表单字段下方有灰色说明，可对照填写。',
        '填写应用/任务名称、应用编码、负责人等应用信息。',
        '选择后端工程类型（Maven JAR、Java WAR/Tomcat、Gradle、Node、Go、Python、Docker 等）与端侧工程类型（无端侧 / Web / uni-app / 微信小程序）。',
        '选择目标主机（须已启用且连通）；填写后端工作区（及可选端侧工程目录）。',
        '按需覆盖后端/端侧验证命令；Docker 可填 Compose 文件与镜像名；Web 可填静态部署目录与本地产物目录。',
        '选择修复引擎：Cursor / CodeX / Claude，或 auto（失败自动切换）。',
        '配置候选部署目录并识别绑定制品；选择是否验证通过后自动部署。',
        '配置巡检：日志模板、周期、时间窗、远端日志路径（须在主机白名单内）等。',
        '保存后建议先不启用，确认配置无误再打开「启用」开关。',
    ])
    add_para(doc, '工程画像与部署提示：', bold=True)
    add_bullets(doc, [
        'JAVA_WAR：工作区为含 pom.xml 的 WAR 工程；部署目标指向 Tomcat webapps；端侧一般选「无端侧」。',
        'Web(npm)：验证通过后可将 dist 同步到目标主机静态目录；留空静态目录则只验证不部署端侧。',
        '微信小程序 / uni-app：以验证为主，平台不自动发布到微信侧。',
        '路径替换类后端（Node/Go/Python/Gradle 等）：按绑定部署目录替换产物并执行重启脚本。',
    ])

    add_heading_cn(doc, '5.2 日常巡检操作', 2)
    add_bullets(doc, [
        '启用：任务按周期自动拉日志；错误达到阈值时自动派生故障事件。',
        '停用：维护窗口、变更窗口或暂不需要自治时，关闭启用开关。',
        '查询：可用关键词、状态、优先级过滤任务列表。',
        '详情：查看任务配置、关联故障与处理进展。',
        '编辑：调整工作区、引擎、周期等；重大变更建议先停用再改。',
        '删除：仅在确认无在途故障/审批时删除，避免中断处置链。',
    ])

    add_heading_cn(doc, '5.3 任务状态理解', 2)
    add_table(
        doc,
        ['关注项', '含义', '运维建议'],
        [
            ('任务状态', '任务本身运行/停用等状态', '异常时先查主机连通与日志白名单'),
            ('故障状态', '由巡检触发的故障处置阶段', '进入详情跟进 Agent 与审批'),
            ('修复引擎', '实际用于本地修码的 CLI', '本机未安装对应 CLI 时勿启用'),
            ('定时拉日志', '是否按周期巡检', '演示/压测后及时关闭，避免无效消耗'),
        ],
    )

    # 6
    add_heading_cn(doc, '六、故障处置流程（运维人员）', 1)
    add_para(doc, '标准闭环如下（系统自动推进 + 人工审批节点）：')
    add_numbered(doc, [
        '巡检拉取日志并判定异常。',
        '派生故障事件，多 Agent 进行证据收集与根因分析。',
        '生成补丁预览（含补丁指纹），进入可审阅状态。',
        '审批人员批准本地修复后，调用所选 CLI 在工作区改码并产出 Diff。',
        '按工程画像执行验证（Maven/Gradle/npm 等）；验证须与真实补丁差异一致。',
        '构建对应制品（JAR/WAR/目录产物等），提供安全下载或进入部署准备。',
        '部署前再次审批/确认；按画像执行路径替换、WAR、Docker 或 Web 静态同步，完成后结案；可选沉淀知识库。',
    ])

    add_heading_cn(doc, '6.1 运维人员跟进清单', 2)
    add_bullets(doc, [
        '在故障详情或执行记录中核对：证据是否充分、根因是否合理、拟修改范围是否越界。',
        '审阅 Diff：确认仅改预期文件，无无关大面积变更。',
        '验证失败时：查看构建/测试日志，判断是环境问题还是补丁问题；必要时驳回并重新规划。',
        '需要人工介入时：可使用 SSH/数据库/Git 客户端协助排查，但不要绕过审批私自上生产。',
    ])

    add_heading_cn(doc, '6.2 执行记录与审计', 2)
    add_bullets(doc, [
        '执行记录用于回放 Agent 步骤、工具调用与关键输出。',
        '处置争议或复盘时，以执行记录 + Diff + 审批记录为准。',
        '勿在处置过程中手动大幅修改同一工作区，以免 Diff/验证不一致。',
    ])

    # 7
    add_heading_cn(doc, '七、审批与治理（管理人员）', 1)
    add_heading_cn(doc, '7.1 必须审批的节点', 2)
    add_table(
        doc,
        ['节点', '审批含义', '通过后系统行为', '驳回建议'],
        [
            ('本地修复前', '是否允许 CLI 改代码', '在工作区执行受控修码', '根因不清/范围过大则驳回'),
            ('部署前', '是否允许进入部署确认', '进入受控部署/完结流程', '验证未过或变更窗口不符则驳回'),
        ],
    )

    add_heading_cn(doc, '7.2 审批操作要点', 2)
    add_bullets(doc, [
        '审批前必看：故障摘要、根因结论、补丁预览/Diff、影响范围、关联主机与工作区。',
        '通过：确认变更可接受且处于允许变更窗口。',
        '驳回：填写清晰原因，便于运维/Agent 调整方案。',
        '重复提交：若阶段已推进，系统可能幂等处理；刷新详情以当前状态为准。',
    ])

    add_heading_cn(doc, '7.3 授权策略与工具管理', 2)
    add_bullets(doc, [
        '在「授权策略」中按角色收敛可调用工具与资源范围。',
        '在「工具管理」中确认诊断类、修复类、验证类工具启用状态符合环境要求。',
        '生产环境遵循「默认拒绝、按需开放」，定期复核策略。',
    ])

    add_heading_cn(doc, '7.4 管理人员日常治理', 2)
    add_bullets(doc, [
        '每周抽查：高优先级故障的 Diff 质量与审批时效。',
        '关注失败率高的修复引擎/主机，安排环境治理。',
        '推动高频故障沉淀至知识库，减少重复研判。',
        '变更窗口外原则上不批准部署类操作。',
    ])

    # 8
    add_heading_cn(doc, '八、知识库使用', 1)
    add_bullets(doc, [
        '处置前：可将历史案例作为根因与方案线索（系统会在可用时引用）。',
        '处置后：将有效根因、补丁要点、规避措施沉淀为知识条目，便于复用。',
        '知识库不可用时，主流程应可降级继续，不得以此阻断紧急处置（按制度执行）。',
        '录入规范：标题清晰、环境标签完整、避免写入明文密码与敏感主机信息。',
    ])
    add_heading_cn(doc, '8.1 文档导入与 AI 预生成', 2)
    add_numbered(doc, [
        '进入「知识库」→「导入/AI 生成」或「新增知识」。',
        '上传文档（推荐 txt/md/docx/log/json/yaml；PDF 请先另存为文本或 docx）和/或填写描述。',
        '点击「AI 预生成」：系统调用模型整理为标题、标签与结构化正文草稿（不落库）。',
        '人工核对、删改敏感信息后点击「保存」入库。',
        '若模型暂不可用，系统会按文档本地整理草稿，仍须人工确认后再保存。',
    ])
    add_para(doc, '注意：预生成仅辅助录入；Agent 召回以已保存且启用的条目为准。', size=10, color=GRAY)

    # 9
    add_heading_cn(doc, '九、人工运维客户端', 1)
    add_heading_cn(doc, '9.1 SSH 客户端', 2)
    add_bullets(doc, [
        '用于人工登录排查；与 Agent 只读诊断互补，不替代审批流程。',
        '仅连接已纳管且授权的主机；遵守操作审计与变更制度。',
    ])
    add_heading_cn(doc, '9.2 数据库终端 / Git 客户端', 2)
    add_bullets(doc, [
        '数据库终端：排查数据相关问题时使用，严禁无审批的数据变更。',
        'Git 客户端：查看分支/提交辅助判断；生产发布仍走既定发布流程。',
    ])

    # 10
    add_heading_cn(doc, '十、修复引擎与制品说明', 1)
    add_heading_cn(doc, '10.1 修复引擎', 2)
    add_table(
        doc,
        ['选项', '说明', '适用场景'],
        [
            ('Cursor CLI', '使用 Cursor 命令行修码', '团队已统一 Cursor 环境'),
            ('CodeX CLI', '使用 CodeX 命令行修码', '已安装并完成授权'),
            ('Claude CLI', '使用 Claude Code 命令行修码', '已安装并完成授权'),
            ('auto', '按策略失败切换引擎', '希望提高成功率，可接受更长耗时'),
        ],
    )
    add_para(doc, '运维要求：任务启用前确认对应 CLI 在运行环境可用；否则巡检发现故障后修码阶段会失败。')

    add_heading_cn(doc, '10.2 验证与制品（多栈）', 2)
    add_bullets(doc, [
        '验证命令默认随 backendKind / clientKind 选择；可在任务中覆盖（端侧仅允许 npm/pnpm/yarn）。',
        'Java（Maven/WAR）构建依赖 JDK 17+；若本机默认 Java 过旧，由管理员配置 ARTIFACT_JAVA_HOME 等。',
        'Gradle / Node / Go / Python / Docker 按各自工具链准备；Docker 需本机或构建机可用 docker/compose。',
        '制品通过平台安全通道处理；自动部署须在审批门禁内执行，禁止私下覆盖生产包绕过系统。',
    ])
    add_heading_cn(doc, '10.3 Agent 环境检测', 2)
    add_bullets(doc, [
        '管理员进入「Agent 管理」，可并排查看 CodeX / Cursor / Claude 的安装与登录状态。',
        'CodeX 支持平台内安装与登录；Cursor / Claude 需本机按官方指引安装后检测与登录。',
        '任务启用前至少保证所选引擎（或 auto 候选引擎）环境可用。',
    ])

    # 11
    add_heading_cn(doc, '十一、日常运行检查（建议）', 1)
    add_table(
        doc,
        ['频率', '检查项', '责任角色'],
        [
            ('每日', '启用中的运维任务是否正常巡检；失败任务是否有告警/记录', '运维人员'),
            ('每日', '待审批事项是否积压；高优先级故障是否超时', '审批/管理人员'),
            ('每周', '主机连通性抽检；凭据是否临期或需轮换', '资源管理员'),
            ('每周', '执行失败原因归类（引擎/权限/工作区/构建）', '运维负责人'),
            ('每月', '授权策略复核；无用任务停用；知识库质量抽检', '系统/运维管理人员'),
        ],
    )

    # 12
    add_heading_cn(doc, '十二、常见问题排查', 1)
    faqs = [
        ('任务启用后一直无故障', '检查目标应用是否确有 ERROR；主机是否启用；日志白名单是否可读；连通性是否通过；巡检周期是否过长。'),
        ('拉日志失败', '复查 SSH 凭据、指纹、白名单路径格式；在目标机确认路径权限；重新探测连通性。'),
        ('新建凭据在主机页不可选', '重新打开主机新增/编辑窗口刷新选项；确认凭据类型匹配 SSH。'),
        ('审批提示不能执行此操作', '故障阶段可能已前进，刷新详情；按当前阶段操作，勿重复提交过期动作。'),
        ('CLI 修码失败', '确认引擎安装与登录授权；工作区路径正确且含工程文件；可换引擎或联系管理员检查策略。'),
        ('验证/制品构建失败', '核对工程画像与本机工具链（JDK/Maven/Node 等）；查看构建日志；清理后重试；仍失败则升级给研发协同。'),
        ('自动部署未生效', '确认已绑定部署目标与制品；检查 autoDeploy 与主机连通；WAR 确认 webapps 路径；Web 确认静态目录已填。'),
        ('Diff 与预期不符', '检查是否有人工同时改动工作区；以执行记录中的补丁指纹与验证结果为准，必要时驳回重来。'),
        ('无菜单权限', '联系系统管理员按角色授权，勿共用管理员账号。'),
    ]
    for q, a in faqs:
        add_para(doc, f'问题：{q}', bold=True)
        add_para(doc, f'处理：{a}', color=GRAY)

    # 13
    add_heading_cn(doc, '十三、安全规范', 1)
    add_bullets(doc, [
        '生产与测试环境严格隔离；测试任务不得指向生产主机或生产工作区。',
        '禁止为图方便关闭审批、放开任意远程命令或扩大日志白名单至敏感目录。',
        '补丁与制品须在授权变更窗口内实施；重大变更双人复核。',
        '账号一人一号；定期审查管理员与审批人名单。',
        '对外材料脱敏：隐藏真实 IP、账号、密钥、内部路径中的敏感段。',
        '发现异常自治行为（越权改文件、异常部署）立即停用相关任务并上报。',
    ])

    # 14
    add_heading_cn(doc, '十四、附录', 1)
    add_heading_cn(doc, '14.1 术语表', 2)
    add_table(
        doc,
        ['术语', '含义'],
        [
            ('运维任务', '绑定主机、日志策略、工作区、工程画像与修复引擎的巡检单元'),
            ('工程画像', 'backendKind/clientKind，决定验证命令、制品发现与部署方式'),
            ('故障事件', '由异常判定触发的一次处置实例'),
            ('补丁指纹', '与补丁内容绑定的校验标识，用于验证/部署防偷换'),
            ('修复引擎', '执行本地改码的 Coding CLI（Cursor/CodeX/Claude）'),
            ('混合部署', '后端与端侧可按画像分别验证/部署（如 JAR + Web 静态）'),
            ('人在回路', '关键有副作用步骤必须人工批准'),
            ('Checkpoint', 'Agent 执行检查点，支持中断后恢复'),
        ],
    )

    add_heading_cn(doc, '14.2 推荐操作顺序（新环境上线）', 2)
    add_numbered(doc, [
        '管理员创建角色权限与审批人。',
        '资源管理员录入凭据并纳管主机（指纹 + 日志白名单 + 探测 + 启用）。',
        '运维人员创建运维任务（先不启用），核对工作区与修复引擎。',
        '在测试应用上试跑：启任务 → 出故障 → 审批 → 修码 → 验证 → 完结。',
        '管理人员复核一次完整审计链后，再逐步扩大任务范围。',
    ])

    add_para(doc, '')
    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = end.add_run('— 全文完 —')
    set_run_font(r, size=10, color=GRAY)

    tip = doc.add_paragraph()
    tip.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tip.add_run('如界面字段随版本变更，以实际系统页面与管理员通知为准。')
    set_run_font(r, size=9, color=GRAY)

    root = Path(__file__).resolve().parents[1]
    out = root / '智维超级员工-运维与管理人员操作手册.docx'
    try:
        doc.save(str(out))
    except PermissionError:
        out = root / '智维超级员工-运维与管理人员操作手册-v2.docx'
        doc.save(str(out))
    print(out)


if __name__ == '__main__':
    main()
