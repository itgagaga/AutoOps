# -*- coding: utf-8 -*-
"""生成《智维超级员工-成果汇报》Word，预留功能截图位供人工粘贴。"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TEAL = RGBColor(0x0B, 0x72, 0x85)
TEAL_DARK = RGBColor(0x0B, 0x4F, 0x5C)
SLATE = RGBColor(0x1F, 0x2A, 0x37)
GRAY = RGBColor(0x52, 0x60, 0x6D)
PLACEHOLDER = RGBColor(0x7A, 0x8B, 0x99)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run_font(run, size=11, bold=False, color=SLATE, name='微软雅黑'):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)


def add_heading_cn(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(
            run,
            size={1: 18, 2: 14, 3: 12}.get(level, 11),
            bold=True,
            color=TEAL_DARK if level == 1 else TEAL,
        )
    return p


def add_para(doc, text, size=11, bold=False, color=SLATE, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_bullets(doc, items, size=11):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.clear()
        run = p.add_run(item)
        set_run_font(run, size=size, color=SLATE)
        p.paragraph_format.space_after = Pt(3)


def set_cell_text(cell, text, bold=False, color=SLATE, size=10, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def shade_cell(cell, fill='F5F8FA'):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    shd.set(qn('w:val'), 'clear')
    tc_pr.append(shd)


def set_row_height(row, cm_height=6.2):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    tr_height = OxmlElement('w:trHeight')
    tr_height.set(qn('w:val'), str(int(Cm(cm_height).twips)))
    tr_height.set(qn('w:hRule'), 'atLeast')
    tr_pr.append(tr_height)


def add_screenshot_slot(doc, fig_id: str, title: str, hint: str, height_cm: float = 6.5):
    """插入可人工替换的截图占位框。"""
    caption = doc.add_paragraph()
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(4)
    run = caption.add_run(f'【截图位 {fig_id}】{title}')
    set_run_font(run, size=11, bold=True, color=TEAL_DARK)

    tip = doc.add_paragraph()
    tip.paragraph_format.space_after = Pt(4)
    run = tip.add_run(f'建议截取：{hint}')
    set_run_font(run, size=9, color=GRAY)

    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    cell = table.rows[0].cells[0]
    shade_cell(cell, 'EEF6F8')
    set_row_height(table.rows[0], height_cm)

    cell.text = ''
    lines = [
        f'图 {fig_id}',
        title,
        '',
        '（请在此处粘贴截图）',
        'Word：插入 → 图片 → 本设备',
        '建议宽度约 14–16 cm，保留界面可读性',
    ]
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2 if i else 18)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line)
        set_run_font(
            run,
            size=12 if i < 2 else 10,
            bold=i < 2,
            color=TEAL_DARK if i < 2 else PLACEHOLDER,
        )

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(2)
    note.paragraph_format.space_after = Pt(10)
    run = note.add_run(f'图 {fig_id}　{title}')
    set_run_font(run, size=9, color=GRAY)
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


def add_table(doc, headers, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, size=10)
        shade_cell(table.rows[0].cells[i], '0B7285')
        for p in table.rows[0].cells[i].paragraphs:
            for run in p.runs:
                run.font.color.rgb = WHITE
                run.font.bold = True
    for r_i, row in enumerate(rows, start=1):
        for c_i, val in enumerate(row):
            set_cell_text(table.rows[r_i].cells[c_i], val, size=10)
    doc.add_paragraph()
    return table


# (fig_id, section_title, feature_bullets, screenshot_title, screenshot_hint)
FEATURE_BLOCKS = [
    (
        '1-1',
        '登录与安全访问',
        [
            '分屏品牌登录页，账号密码 + 滑块安全验证。',
            '登录后按角色加载菜单与数据权限。',
        ],
        '登录页全貌',
        '浏览器全屏登录页（含左侧品牌区与右侧登录表单、滑块验证）',
    ),
    (
        '2-1',
        '运维工作台总览',
        [
            '首页/仪表盘展示任务、故障、审批等关键运行态势。',
            '便于值班人员快速进入待办与异常任务。',
        ],
        '工作台 / 仪表盘',
        '登录后首页或仪表盘整页（含关键指标与入口）',
    ),
    (
        '3-1',
        '运维任务与日志巡检',
        [
            '配置主机、工作区、日志模板与周期拉取。',
            '支持异常判定后自动派生故障事件并触发 Agent 分析。',
        ],
        '运维任务列表',
        '运维任务列表页，可见任务状态、最近拉取摘要',
    ),
    (
        '3-2',
        '运维任务详情配置',
        [
            '任务级工程画像：后端栈（Maven/WAR/Gradle/Node/Go/Python/Docker）与端侧类型。',
            '可配置修复引擎偏好、自动部署与验证命令。',
        ],
        '运维任务详情 / 编辑表单',
        '打开某一任务详情或编辑弹窗，露出工程画像与部署相关字段',
    ),
    (
        '4-1',
        '多 Agent 故障分析',
        [
            '协调智能体委派证据收集、根因分析、修复规划等子智能体。',
            '时间线展示分析过程与工具调用结果。',
        ],
        '故障事件时间线 / Agent 分析过程',
        '故障详情页时间线，含多 Agent 协作步骤',
    ),
    (
        '5-1',
        '本地修复引擎（CodeX / Cursor / Claude）',
        [
            '支持 CodeX、Cursor、Claude Code 三引擎及 auto 偏好。',
            'Agent 管理页可并排检测 CLI 环境与登录状态。',
        ],
        '修复引擎环境检测',
        'Agent 管理或系统设置中 CodeX / Cursor / Claude 环境状态区',
    ),
    (
        '5-2',
        '代码修复与差异确认',
        [
            '补丁预览 → 本地 CLI 受控修复 → 真实差异确认。',
            '确认后异步构建制品并进入部署/人工验证环节。',
        ],
        '修复确认 / 差异预览',
        '含 Cursor/CodeX 修复结果或等待确认差异的界面',
    ),
    (
        '6-1',
        '验证、审批与部署',
        [
            'TEST_VERIFICATION 校验修复结果；高风险动作走审批。',
            '多栈混合部署：路径替换 / WAR / Docker / Web 静态同步等。',
        ],
        '审批中心或部署结果',
        '待审批列表，或部署成功/一键脚本生成结果页',
    ),
    (
        '7-1',
        '知识库导入与 AI 预生成',
        [
            '支持 txt/md/docx 等文档导入。',
            '可基于文档或描述调用 Agent 预生成草稿，人工确认后入库。',
        ],
        '知识库导入 / AI 预生成',
        '知识库页面：上传区、AI 预生成按钮与草稿编辑区',
    ),
    (
        '8-1',
        'Agent 与权限治理',
        [
            'Agent 清单、执行记录、工具授权与审批策略。',
            '用户 / 角色 / 字典等系统管理能力。',
        ],
        'Agent 管理或权限配置',
        'Agent 管理列表，或某 Agent 的工具/资源权限配置页',
    ),
]


def main():
    root = Path(__file__).resolve().parents[1]
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    # Cover
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(36)
    r = t.add_run('智维超级员工')
    set_run_font(r, size=28, bold=True, color=TEAL_DARK)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run('产品成果汇报')
    set_run_font(r, size=20, bold=True, color=TEAL)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run('面向服务器运维的多 Agent 全流程自治平台')
    set_run_font(r, size=12, color=GRAY)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(16)
    r = meta.add_run('版本：V1.2 ｜ 含功能截图占位 ｜ 内部资料')
    set_run_font(r, size=10, color=GRAY)

    add_para(
        doc,
        '本文档汇总本期产品成果与核心能力说明。文中灰色框为【截图位】，请按图号在对应位置粘贴真实界面截图后用于汇报或验收。'
        '粘贴后可删除框内提示文字，仅保留图片与图题。',
        size=10,
        color=GRAY,
    )

    add_heading_cn(doc, '一、成果摘要', 1)
    add_para(doc, '在「日志巡检 → 多 Agent 分析 → CLI 修码 → 验证 → 审批 → 部署」闭环基础上，本期完成能力收敛与交付固化：')
    add_bullets(doc, [
        '多技术栈工程画像：后端 Maven JAR / Java WAR / Gradle / Node / Go / Python / Docker；端侧 Web(npm)、uni-app、微信原生小程序等。',
        '混合部署：按栈路由路径替换、WAR、Docker Compose、Web 静态同步；小程序仅验证不自动发布。',
        '修复引擎三选一 + auto：CodeX / Cursor / Claude Code；管理页环境检测与登录引导。',
        '知识库：文档导入 + Agent 预生成草稿，人工确认入库。',
        '运维体验：列表默认 20 条、表单字段说明、长弹窗内部滚动、登录页品牌化改版。',
        '并发与达梦兼容加固：会话 CAS、租约释放、拉日志多实例互斥、NCLOB 正文读取等。',
    ])

    add_heading_cn(doc, '二、端到端能力链路', 1)
    add_table(doc, ['阶段', '能力要点', '对应截图'], [
        ('接入与登录', '安全登录、角色权限', '图 1-1'),
        ('巡检配置', '运维任务、日志拉取、工程画像', '图 3-1 / 3-2'),
        ('智能分析', '多 Agent 协同、时间线', '图 4-1'),
        ('自动修复', 'CLI 修码、差异确认', '图 5-1 / 5-2'),
        ('验证部署', '验证、审批、多栈部署', '图 6-1'),
        ('知识沉淀', '导入与 AI 预生成', '图 7-1'),
        ('平台治理', 'Agent / 权限 / 字典', '图 8-1'),
    ])

    add_heading_cn(doc, '三、截图填写说明', 1)
    add_bullets(doc, [
        '分辨率：优先 1920×1080 浏览器窗口；避免只截局部导致字迹不清。',
        '脱敏：隐藏真实主机 IP、账号、密钥、客户内部路径等敏感信息。',
        '命名建议：将原始截图另存为 `screenshots/图X-X-简短标题.png`，便于版本管理。',
        '替换方式：选中占位框内文字 → 删除 →「插入 → 图片」；或直接覆盖单元格内容。',
        '图题已预置在框下方，粘贴图片后请保留「图 X-X　标题」行。',
    ])

    add_heading_cn(doc, '四、功能成果与截图位', 1)
    for index, (fig_id, section_title, bullets, shot_title, hint) in enumerate(FEATURE_BLOCKS, start=1):
        add_heading_cn(doc, f'4.{index} {section_title}', 2)
        add_para(doc, '能力说明：', size=10, bold=True, color=TEAL, space_after=2)
        add_bullets(doc, bullets)
        add_screenshot_slot(doc, fig_id, shot_title, hint)

    add_heading_cn(doc, '五、截图对照清单（勾选用）', 1)
    add_para(doc, '人工截图完成后，可在下表「完成」列打勾，便于移交核对。')
    checklist_rows = []
    for fig_id, section_title, _bullets, shot_title, hint in FEATURE_BLOCKS:
        checklist_rows.append((fig_id, shot_title, section_title, '☐'))
    add_table(doc, ['图号', '截图标题', '所属能力', '完成'], checklist_rows)

    add_heading_cn(doc, '六、附录：建议补充截图（可选）', 1)
    add_para(doc, '若汇报篇幅允许，可按需追加以下截图位（本文未预置框，可复制第四节样式自行增加）：')
    add_bullets(doc, [
        '图 A-1：资源中心主机列表 / SSH 连通',
        '图 A-2：字典管理或菜单权限',
        '图 A-3：知识库条目详情',
        '图 A-4：多目标部署路径配置',
        '图 A-5：故障完结与知识沉淀提示',
    ])

    add_heading_cn(doc, '七、版本与编制说明', 1)
    add_table(doc, ['项', '内容'], [
        ('产品版本', 'V1.2'),
        ('文档用途', '成果汇报 / 验收附图底稿'),
        ('截图状态', '占位待填（人工粘贴）'),
        ('关联交付', 'delivery/ 交付包、操作手册、路演 PPT'),
    ])

    out = root / '智维超级员工-成果汇报.docx'
    try:
        doc.save(str(out))
    except PermissionError:
        out = root / '智维超级员工-成果汇报-v2.docx'
        doc.save(str(out))
    print(f'已生成: {out}')
    return out


if __name__ == '__main__':
    main()
