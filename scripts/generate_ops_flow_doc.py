# -*- coding: utf-8 -*-
"""生成《智维超级员工-操作流程技术说明》Word，含功能截图占位。"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TEAL = RGBColor(0x0B, 0x72, 0x85)
TEAL_DARK = RGBColor(0x0B, 0x4F, 0x5C)
SLATE = RGBColor(0x1F, 0x2A, 0x37)
GRAY = RGBColor(0x52, 0x60, 0x6D)
PLACEHOLDER = RGBColor(0x7A, 0x8B, 0x99)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run_font(run, size=11, bold=False, color=SLATE, name='微软雅黑'):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)


def add_heading_cn(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(
            run,
            size={1: 18, 2: 14, 3: 12}.get(level, 11),
            bold=True,
            color=TEAL_DARK if level == 1 else TEAL,
        )
    return p


def add_para(doc, text, size=11, bold=False, color=SLATE, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_bullets(doc, items, size=11):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.clear()
        run = p.add_run(item)
        set_run_font(run, size=size, color=SLATE)
        p.paragraph_format.space_after = Pt(3)


def add_numbered(doc, items, size=11):
    for item in items:
        p = doc.add_paragraph(style='List Number')
        p.clear()
        run = p.add_run(item)
        set_run_font(run, size=size, color=SLATE)
        p.paragraph_format.space_after = Pt(4)


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size=9, color=SLATE, name='Consolas')
    run.font.name = 'Consolas'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), 'F5F8FA')
    shd.set(qn('w:val'), 'clear')
    p._p.get_or_add_pPr().append(shd)
    return p


def set_cell_text(cell, text, bold=False, color=SLATE, size=10, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def shade_cell(cell, fill='F5F8FA'):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    shd.set(qn('w:val'), 'clear')
    tc_pr.append(shd)


def shade_header_row(row):
    for cell in row.cells:
        shade_cell(cell, '0B7285')
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = WHITE
                run.font.bold = True


def set_row_height(row, cm_height=6.2):
    tr = row._tr
    tr_pr = tr.get_or_add_trPr()
    tr_height = OxmlElement('w:trHeight')
    tr_height.set(qn('w:val'), str(int(Cm(cm_height).twips)))
    tr_height.set(qn('w:hRule'), 'atLeast')
    tr_pr.append(tr_height)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, size=10)
    shade_header_row(table.rows[0])
    for r_i, row in enumerate(rows, start=1):
        for c_i, val in enumerate(row):
            set_cell_text(table.rows[r_i].cells[c_i], val, size=10)
    doc.add_paragraph()
    return table


def add_screenshot_slot(doc, fig_id, title, hint, height_cm=6.5):
    caption = doc.add_paragraph()
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(4)
    run = caption.add_run(f'【截图位 {fig_id}】{title}')
    set_run_font(run, size=11, bold=True, color=TEAL_DARK)

    tip = doc.add_paragraph()
    tip.paragraph_format.space_after = Pt(4)
    run = tip.add_run(f'建议截取：{hint}')
    set_run_font(run, size=9, color=GRAY)

    table = doc.add_table(rows=1, cols=1)
    table.style = 'Table Grid'
    cell = table.rows[0].cells[0]
    shade_cell(cell, 'EEF6F8')
    set_row_height(table.rows[0], height_cm)

    cell.text = ''
    lines = [
        f'图 {fig_id}',
        title,
        '',
        '（请在此处粘贴截图）',
        'Word：插入 → 图片 → 本设备',
        '建议宽度约 14–16 cm，隐藏真实 IP / 账号 / 密钥',
    ]
    for i, line in enumerate(lines):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2 if i else 18)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line)
        set_run_font(
            run,
            size=12 if i < 2 else 10,
            bold=i < 2,
            color=TEAL_DARK if i < 2 else PLACEHOLDER,
        )

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(2)
    note.paragraph_format.space_after = Pt(10)
    run = note.add_run(f'图 {fig_id}　{title}')
    set_run_font(run, size=9, color=GRAY)
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


SHOTS = [
    ('1-1', '登录页', '浏览器全屏登录页，含验证码/滑块'),
    ('2-1', '操作指南首页', '登录后 HOME 页，含四阶段路线与「进入运维工作台」'),
    ('3-1', '本地修复引擎向导', 'CodeX / Cursor / Claude 检测、安装或登录状态为可用'),
    ('4-1', '凭据管理', '凭据列表或新增表单（密码打码）'),
    ('4-2', '主机管理', '主机编辑：指纹、日志白名单、连通性探测结果'),
    ('5-1', 'Agent 授权', '协调 Agent / 专业 Agent 工具与资源授权'),
    ('5-2', '知识库', '知识列表或导入/AI 预生成草稿页'),
    ('6-1', '新建运维任务', '任务表单：工程画像、工作区、主机、巡检周期'),
    ('6-2', '运维工作台列表', '已启用任务、最近拉日志摘要、关联故障'),
    ('7-1', '故障详情时间线', '多 Agent 步骤：取证、根因、补丁预览'),
    ('7-2', '补丁差异审阅', 'RepairDiff / 变更文件列表'),
    ('8-1', '修复审批', '待审批「本地修复」记录，含范围与风险'),
    ('8-2', '验证结果', 'TEST_VERIFICATION 通过或失败日志摘要'),
    ('9-1', '部署审批与结果', '二次审批、部署目标或脚本下载'),
    ('10-1', 'Agent 执行记录', '会话、工具调用、耗时'),
]


def main():
    root = Path(__file__).resolve().parents[1]
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(28)
    r = t.add_run('智维超级员工')
    set_run_font(r, size=28, bold=True, color=TEAL_DARK)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run('操作流程技术说明')
    set_run_font(r, size=20, bold=True, color=TEAL)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run('从环境准备到巡检、分析、修码、验证、审批部署的端到端技术流程')
    set_run_font(r, size=12, color=GRAY)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(12)
    r = meta.add_run('版本：V1.2 ｜ 适用：实施、运维、研发联调 ｜ 含功能截图占位')
    set_run_font(r, size=10, color=GRAY)

    add_para(
        doc,
        '本文按真实系统行为描述操作顺序、页面入口、工具门禁与故障状态机。'
        '灰色框为【截图位】，按图号粘贴脱敏后的界面即可用于培训、验收或移交。'
        '日常按钮级说明另见《运维与管理人员操作手册》；本文偏技术闭环与约束。',
        size=10,
        color=GRAY,
    )

    add_heading_cn(doc, '一、文档范围与阅读对象', 1)
    add_table(doc, ['读者', '关注内容'], [
        ('实施 / 交付', '准备顺序、工作区、CLI、主机指纹、SQL 已执行'),
        ('运维值班', '任务启停、故障跟进、Diff 审阅、知识回填'),
        ('审批人', '两次门禁看什么、驳回后系统停在哪'),
        ('研发联调', '工具码、状态迁移、为何乱序会被拒绝'),
    ])
    add_para(doc, '闭环总览（业务顺序，不可跳步）：', bold=True)
    add_code(
        doc,
        '环境准备 → 资源纳管 → Agent 授权\n'
        '→ 创建并启用运维任务（周期拉日志）\n'
        '→ 异常建事件 → LOG_QUERY → 根因 / PATCH_PREVIEW\n'
        '→ 审批① CODEX_LOCAL_REPAIR（本地 CLI 修码）\n'
        '→ TEST_VERIFICATION（与真实 diff 绑定）\n'
        '→ 审批② DEPLOY_PATCH → 完结 → 知识沉淀',
    )

    add_heading_cn(doc, '二、系统角色与页面入口', 1)
    add_table(doc, ['能力', '菜单 code / 入口', '流程中的位置'], [
        ('操作指南', 'HOME', '四阶段路线，可跳转各准备项'),
        ('修复引擎向导', 'CODEX_INSTALL_GUIDE（权限 AGENT_MANAGEMENT）', '准备：CLI 可用'),
        ('主机', 'RESOURCE_HOST', '准备：SSH 与日志白名单'),
        ('凭据', 'RESOURCE_CREDENTIAL', '准备：加密凭据'),
        ('运维工作台', 'OPS_DASHBOARD', '任务、巡检、故障主战场'),
        ('Agent 授权', 'AGENT_AUTHORIZATION', '工具与资源范围'),
        ('知识库', 'OPS_KNOWLEDGE', '分析上下文 / 事后沉淀'),
        ('执行记录', 'AGENT_EXECUTION', '会话与工具调用回放'),
        ('审批', '事件详情 / 待办（OPS_APPROVAL 默认隐藏）', '修码前、部署前'),
        ('SSH / DB / Git', 'OPS_SSH_CLIENT 等', '人工排障，不替代审批'),
    ])
    add_screenshot_slot(doc, '1-1', '登录页', '浏览器全屏登录页，含验证码或滑块')
    add_screenshot_slot(doc, '2-1', '操作指南首页', 'HOME：四阶段路线与进入工作台按钮')

    add_heading_cn(doc, '三、阶段 A：环境与资源准备（一次性）', 1)
    add_para(doc, '未完成准备时不要启用巡检。缺 CLI、缺指纹或日志路径不在白名单，后续阶段会失败。')

    add_heading_cn(doc, '3.1 本地修复引擎', 2)
    add_numbered(doc, [
        '进入「本地修复引擎向导」（或 Agent 管理中的环境检测）。',
        '选择默认引擎：codex / cursor / claude；任务级可覆盖。',
        '检测 CLI：codex --version、agent、claude。不可用则按向导安装并登录。',
        '确认页面显示「环境可用」后再建任务。',
    ])
    add_para(doc, '技术约束：CLI 只在含工程标识的本地工作区执行（Maven 须含 pom.xml）；远程证据只走 SSH 只读模板。', size=10, color=GRAY)
    add_screenshot_slot(doc, '3-1', '本地修复引擎向导', '三引擎并排检测，至少一项为可用')

    add_heading_cn(doc, '3.2 凭据与主机', 2)
    add_numbered(doc, [
        '凭据管理：新增 SSH（及按需数据库/Git）凭据；密钥入库加密，页面不回显。',
        '主机管理：IP、端口、绑定凭据、操作系统、环境。',
        '填写 SHA-256 主机密钥指纹（ssh-keyscan | ssh-keygen -lf - -E sha256）。',
        '配置允许读取的日志绝对路径（文件或目录）；禁止 .. 与相对路径。',
        '先保存为未启用 → 探测连通性 → 通过后再启用。',
    ])
    add_code(doc, 'ssh-keyscan -p <端口> <主机IP> 2>$null | ssh-keygen -lf - -E sha256')
    add_screenshot_slot(doc, '4-1', '凭据管理', '凭据列表或新增表单，密码区域打码')
    add_screenshot_slot(doc, '4-2', '主机管理', '指纹列表、日志白名单、连通性成功')

    add_heading_cn(doc, '3.3 Agent 授权与知识库', 2)
    add_numbered(doc, [
        '确认协调 Agent 与四个专业 Agent 在线（默认同进程，无需起 8101–8104）。',
        '授权诊断类工具（日志查询等）及修码/验证/部署工具；DEPLOY_PATCH 生产默认收紧。',
        '为主机授予 HOST 读取范围；凭据访问策略按环境打开。',
        '知识库可先导入制度/历史案例；AI 预生成只出草稿，保存后才参与分析召回。',
    ])
    add_screenshot_slot(doc, '5-1', 'Agent 授权', '工具授权与资源范围配置')
    add_screenshot_slot(doc, '5-2', '知识库', '导入区或 AI 预生成草稿（未点保存前）')

    add_heading_cn(doc, '四、阶段 B：创建并启用运维任务', 1)
    add_numbered(doc, [
        '运维工作台 → 新增运维任务。',
        '填写标题、应用编码、优先级（P0–P3）。',
        'backendKind：MAVEN / JAVA_WAR / GRADLE / NODE / GO / PYTHON / DOCKER。',
        'clientKind：NONE / WEB_NPM / UNI_APP / MP_WEIXIN_NATIVE。',
        '选择已启用主机；填写后端工作区（Git 仓库根，须能 git diff）；可选端侧目录。',
        '选择修复引擎或 auto；配置验证命令（可覆盖默认）；绑定部署目录/制品。',
        '巡检：日志模板（JOURNAL_LOGS / APP_LOG_FILE / JAVA_APPLICATION_LOGS 等）、周期、时间窗、远端路径（须在白名单）。',
        '建议先保存不启用，核对后再打开启用开关。启用后按周期拉日志，异常达阈值派生故障事件。',
    ])
    add_screenshot_slot(doc, '6-1', '新建运维任务', '表单露出工程画像、工作区、主机、巡检字段')
    add_screenshot_slot(doc, '6-2', '运维工作台列表', '已启用任务、最近拉取、关联故障入口')

    add_heading_cn(doc, '五、阶段 C：自动巡检到多 Agent 分析', 1)
    add_para(doc, '调度线程只提交拉日志任务，拉取与分析在工作线程执行。异常判定见配置：最少 ERROR 条数、增长比等。')
    add_numbered(doc, [
        '到周期后 SSH 按模板拉取，日志落盘工作区 .zhiwei-logs/（不入库 Git）。',
        '达到异常条件则创建故障事件，状态 NEW → RUNNING。',
        '协调 Agent 调用日志提取：TRIAGED。',
        '根因分析：DIAGNOSED。',
        '修复规划 PATCH_PREVIEW，写入补丁指纹与允许修改文件：PATCH_READY。',
        '值班在故障详情时间线核对证据与方案；范围不对则不要批修码。',
    ])
    add_para(doc, '只读 SSH 模板（模型不能下发任意 Shell）：HOST_HEALTH、SERVICE_STATUS、JOURNAL_LOGS、DISK_USAGE、JAVA_APPLICATION_LOGS、APP_LOG_FILE。')
    add_screenshot_slot(doc, '7-1', '故障详情时间线', '取证 / 根因 / 补丁预览步骤可见')

    add_heading_cn(doc, '六、阶段 D：审批修码与差异确认', 1)
    add_para(doc, '门禁 1：CODEX_LOCAL_REPAIR（名称沿用，实际引擎可以是 Cursor / Claude）。')
    add_numbered(doc, [
        '系统在 PATCH_READY 且已有证据、指纹未变、允许文件非空时，发起修码审批（ASK）。',
        '审批人查看根因、文件白名单、工作区路径；通过则启动 CLI（沙箱/工作区写入，禁止 commit/push/deploy）。',
        '执行后读取 git diff；越界则失败并作废后续测试候选。',
        '成功则写入实际 diff，清除旧 verificationPassed；界面弹出差异供「接受 / 撤销」。',
        '接受后进入验证；撤销则回到可重新规划，不进入部署。',
    ])
    add_screenshot_slot(doc, '8-1', '修复审批', '待办中的本地修复审批，含范围说明')
    add_screenshot_slot(doc, '7-2', '补丁差异审阅', '文件列表与 hunk，可接受或撤销')

    add_heading_cn(doc, '七、阶段 E：验证、二次审批与部署', 1)
    add_para(doc, '门禁 2：验证必须使用修码返回的实际 diff；部署指纹须与验证指纹、已批准部署键一致。')
    add_numbered(doc, [
        'TEST_VERIFICATION：按 backendKind/clientKind 执行验证命令（Java 常用 mvn test / verify）。',
        '通过后状态趋向 APPROVAL_PENDING；可构建制品（JAR/WAR/目录），JDK 须 17+。',
        '部署前再次人工审批 DEPLOY_PATCH。',
        '按画像执行：路径替换、WAR→webapps、Docker Compose、Web 静态同步；小程序只验证不自动发布。',
        '成功 → RESOLVED；失败 → FAILED。可选将根因与处置写入知识库。',
    ])
    add_screenshot_slot(doc, '8-2', '验证结果', '验证通过摘要或失败日志关键段')
    add_screenshot_slot(doc, '9-1', '部署审批与结果', '二次审批或部署成功/脚本下载')

    add_heading_cn(doc, '八、阶段 F：回放与人工排障', 1)
    add_bullets(doc, [
        '执行记录：会话、工具码、耗时、失败步骤；复盘以记录 + Diff + 审批为准。',
        'SSH/DB/Git 客户端仅用于人工协助，禁止绕过状态机改生产。',
        '处置过程中勿手工大改同一工作区，以免 diff 与验证不一致。',
    ])
    add_screenshot_slot(doc, '10-1', 'Agent 执行记录', '工具调用列表与关键输出')

    add_heading_cn(doc, '九、状态机与工具门禁（实现约束）', 1)
    add_heading_cn(doc, '9.1 故障状态（IncidentStatus）', 2)
    add_table(doc, ['状态', '含义', '允许下一状态（摘要）'], [
        ('NEW', '已建事件未跑自动流', 'RUNNING / CANCELLED'),
        ('RUNNING', '自动处置进行中', 'TRIAGED / FAILED / TIMEOUT / WAITING_INPUT'),
        ('TRIAGED', '日志范围已定', 'DIAGNOSED / FAILED / WAITING_INPUT'),
        ('DIAGNOSED', '根因已出', 'PATCH_READY / FAILED / WAITING_INPUT'),
        ('PATCH_READY', '补丁草案就绪', 'TESTING / FAILED / WAITING_INPUT'),
        ('TESTING', '正在验证', 'APPROVAL_PENDING / FAILED / WAITING_INPUT'),
        ('APPROVAL_PENDING', '待部署审批', 'DEPLOYING / REJECTED / FAILED / WAITING_INPUT'),
        ('DEPLOYING', '正在部署', 'RESOLVED / FAILED / WAITING_INPUT'),
        ('WAITING_INPUT', '等人（含审批）', 'RUNNING / RESOLVED / FAILED / CANCELLED'),
        ('RESOLVED / CANCELLED / TIMEOUT / REJECTED / FAILED', '终态', '不可再进自动主流程'),
    ])
    add_para(
        doc,
        '实现说明：本地修码发生在 PATCH_READY 之后、写入 TESTING 之前；审批等待使用 WAITING_INPUT 或待办动作。'
        '乱序调用（未预览就修码、验证 diff 与实际不符、指纹变了还部署）由 IncidentLifecycleService 拒绝。',
        size=10,
        color=GRAY,
    )

    add_heading_cn(doc, '9.2 工具调用顺序', 2)
    add_table(doc, ['工具码', '作用', '前置', '人'], [
        ('日志/健康类 LOG_QUERY 等', 'SSH 只读取证', '主机授权+指纹+白名单', '否'),
        ('PATCH_PREVIEW', '补丁计划+指纹', '已有证据', '审阅'),
        ('CODEX_LOCAL_REPAIR', '工作区受控改文件', 'PATCH_READY + 审批', '必须批'),
        ('TEST_VERIFICATION', '多栈验证', '实际 diff 一致', '看结果'),
        ('DEPLOY_PATCH', '受控部署', '验证指纹+再审批', '必须批'),
    ])

    add_heading_cn(doc, '十、异常与回退', 1)
    add_table(doc, ['现象', '常见原因', '处理'], [
        ('任务不拉日志', '未启用、主机未通、路径不在白名单、SSH 开关关', '按 3.2 / 阶段 B 排查'),
        ('分析停住', '模型 Key、Agent 未授权、超时租约', '看执行记录；恢复扫描会重试有限次数'),
        ('修码失败', 'CLI 未装、工作区无 Git、越界改文件', '引擎向导；缩小白名单后重新预览'),
        ('验证失败', '补丁不对、JDK 过旧、命令非白名单', '驳回修码；检查 ARTIFACT_JAVA_HOME'),
        ('部署被拒', '未二次审批、指纹不一致、非变更窗口', '按制度驳回或改窗口后再批'),
        ('审批重复点', '状态已前进', '以详情当前态为准，幂等忽略'),
    ])

    add_heading_cn(doc, '十一、演示工程最短路径', 1)
    add_numbered(doc, [
        '启动「测试工程」demo-ops-app，确认 logs/demo-ops-app.log 有 ERROR（或 scripts/trigger-faults.ps1）。',
        '工作区对该目录执行 scripts/init-git.ps1，保证独立 Git 可供 diff。',
        '主机日志白名单指向该日志绝对路径（Windows 含中文路径时建议 WSL 映射）。',
        '任务指向该工作区与主机，短周期、开自动分析。',
        '按阶段 C–E 点两次审批，走完验证；演示部署可用受控命令代替生产发布。',
    ])

    add_heading_cn(doc, '十二、截图填写说明与对照清单', 1)
    add_bullets(doc, [
        '分辨率优先 1920×1080；脱敏 IP、账号、密钥、客户路径。',
        '选中占位框文字删除后「插入 → 图片」；保留框下「图 X-X」题注。',
        '原始文件建议另存 screenshots/图X-X-标题.png。',
    ])
    add_table(doc, ['图号', '标题', '完成'], [(i, t, '☐') for i, t, _h in SHOTS])

    add_heading_cn(doc, '十三、版本与关联文档', 1)
    add_table(doc, ['项', '内容'], [
        ('产品版本', 'V1.2'),
        ('本文定位', '操作流程技术说明（含截图位）'),
        ('关联', '操作手册、路演手册、AI 提效说明、CODEX-CLI-REPAIR.md、HOST-APP-LOG-PATH.md'),
        ('重新生成', 'python .\\scripts\\generate_ops_flow_doc.py'),
    ])

    out = root / '智维超级员工-操作流程技术说明.docx'
    try:
        doc.save(str(out))
    except PermissionError:
        out = root / '智维超级员工-操作流程技术说明-v2.docx'
        doc.save(str(out))
    print(f'已生成: {out}')
    return out


if __name__ == '__main__':
    main()
