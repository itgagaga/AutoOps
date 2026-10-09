# -*- coding: utf-8 -*-
"""Build delivery package: consolidated Dameng SQL + Word checklists."""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SQL_SRC = ROOT / 'super-employee-platform' / 'super-employee-app' / 'docs' / 'sql'
DELIVERY = ROOT / 'delivery'
DELIVERY_SQL = DELIVERY / 'sql' / 'dameng'

TEAL_DARK = RGBColor(0x0B, 0x4F, 0x5C)
TEAL = RGBColor(0x0B, 0x72, 0x85)
SLATE = RGBColor(0x1F, 0x2A, 0x37)
GRAY = RGBColor(0x52, 0x60, 0x6D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# 存量库升级顺序（达梦）。已执行过的脚本务必跳过。
UPGRADE_ORDER = [
    ('01', 'migration-dameng-rbac.sql', 'RBAC 基础'),
    ('02', 'migration-dameng-agent-management.sql', 'Agent/工具管理'),
    ('03', 'migration-dameng-agent-authorization.sql', 'Agent 授权与审批策略表'),
    ('04', 'migration-dameng-resource-center.sql', '资源中心'),
    ('05', 'migration-dameng-operations.sql', '运维操作相关'),
    ('06', 'dict-migration-dameng.sql', '数据字典'),
    ('07', 'ops-task-migration-dameng.sql', '运维任务表'),
    ('08', 'ops-task-workspace-migration-dameng.sql', '工作区字段'),
    ('09', 'ops-task-application-deployment-migration-dameng.sql', '应用/部署字段'),
    ('10', 'ops-task-deployment-paths-migration-dameng.sql', '部署路径列表'),
    ('11', 'ops-task-deployment-jar-migration-dameng.sql', '部署 JAR'),
    ('12', 'ops-task-deployment-targets-migration-dameng.sql', '部署目标'),
    ('13', 'ops-log-schedule-migration-dameng.sql', '日志调度'),
    ('14', 'ops-log-range-migration-dameng.sql', '日志时间窗'),
    ('15', 'ops-task-repair-engine-migration-dameng.sql', '修复引擎字段'),
    ('15b', 'ops-task-frontend-workspace-migration-dameng.sql', '前端工程目录字段'),
    ('15c', 'ops-task-project-profile-migration-dameng.sql', '工程画像字段（后端/端侧类型）'),
    ('15d', 'ops-task-multi-stack-deploy-migration-dameng.sql', '多栈混合部署字段'),
    ('16', 'ops-data-permission-migration-dameng.sql', '运维数据权限'),
    ('17', 'host-log-path-migration-dameng.sql', '主机日志白名单'),
    ('18', 'db-connectivity-migration-dameng.sql', '数据库连通性'),
    ('19', 'db-inline-auth-migration-dameng.sql', '数据库内联认证'),
    ('20', 'workspace-settings-migration-dameng.sql', '工作区设置'),
    ('21', 'incident-large-text-migration-dameng.sql', '故障大文本'),
    ('22', 'react-agent-migration.sql', 'Agent 会话/轮次/工具调用'),
    ('23', 'react-agent-approval-migration.sql', '待审批动作与工具种子'),
    ('24', 'checkpoint-recovery-migration-dameng.sql', 'Checkpoint/租约'),
    ('25', 'ssh-state-machine-migration-dameng.sql', '补丁指纹门禁字段'),
    ('26', 'agent-session-error-clob-migration-dameng.sql', '会话错误 CLOB'),
    ('27', 'agent-i18n-migration-dameng.sql', 'Agent 国际化'),
    ('28', 'codex-repair-migration.sql', 'CODEX_LOCAL_REPAIR 工具'),
    ('29', 'knowledge-base-migration.sql', '知识库表与菜单'),
    ('30', 'knowledge-created-by-migration-dameng.sql', '知识库创建人（若旧库缺列）'),
    ('31', 'ssh-client-menu-dameng.sql', 'SSH 客户端菜单'),
    ('32', 'db-client-menu-dameng.sql', '数据库终端菜单'),
    ('33', 'git-client-menu-dameng.sql', 'Git 客户端菜单'),
    ('34', 'menu-order-migration-dameng.sql', '菜单排序'),
    ('35', 'tool-cleanup-git-enable-dameng.sql', '工具清理/Git 启用'),
]

OPTIONAL_SEED = [
    ('optional', 'resource-demo-seed-dameng.sql', '演示资源种子（仅测试环境）'),
]


def read_sql(name: str) -> str:
    return (SQL_SRC / name).read_text(encoding='utf-8')


def build_install_full() -> str:
    schema = read_sql('schema-dameng.sql')
    if 'repair_engine' not in schema:
        schema = schema.replace(
            '    log_path VARCHAR(1024),\n',
            '    repair_engine VARCHAR(32),\n    log_path VARCHAR(1024),\n',
            1,
        )

    parts = [
        '-- =============================================================================',
        '-- 智维超级员工 · 达梦数据库 · 新环境一键初始化（整合版）',
        '-- 生成方式: python scripts/build_delivery_package.py',
        '-- 适用: 空库/空模式，仅执行一次',
        '-- 注意: 非幂等 DDL 为主；已有库请改用 upgrade 清单，勿直接执行本文件',
        '-- =============================================================================',
        '',
        '-- ---------- A. 基础 schema（含菜单/角色种子）----------',
        schema.rstrip(),
        '',
        '-- ---------- B. Agent 授权相关表 ----------',
        read_sql('migration-dameng-agent-authorization.sql').rstrip(),
        '',
        '-- ---------- C. Agent 运行时表（已合并 Checkpoint 列）----------',
        '''
CREATE TABLE agent_session (
    id VARCHAR(64) PRIMARY KEY,
    incident_id VARCHAR(64) NOT NULL,
    agent_code VARCHAR(64) NOT NULL,
    thread_id VARCHAR(128) NOT NULL UNIQUE,
    status VARCHAR(32) NOT NULL,
    objective VARCHAR(4000) NOT NULL,
    final_summary CLOB,
    error_message CLOB,
    started_at VARCHAR(64),
    finished_at VARCHAR(64),
    created_at VARCHAR(64) NOT NULL,
    updated_at VARCHAR(64) NOT NULL,
    attempt_count INTEGER DEFAULT 0 NOT NULL,
    max_attempts INTEGER DEFAULT 3 NOT NULL,
    next_retry_at VARCHAR(64),
    lease_owner VARCHAR(128),
    lease_until VARCHAR(64),
    last_heartbeat_at VARCHAR(64),
    failure_type VARCHAR(64),
    retryable SMALLINT DEFAULT 0 NOT NULL,
    version BIGINT DEFAULT 0 NOT NULL
);

CREATE INDEX idx_agent_session_recovery
    ON agent_session(status, next_retry_at, lease_until);

CREATE TABLE agent_turn (
    id VARCHAR(64) PRIMARY KEY,
    session_id VARCHAR(64) NOT NULL,
    turn_no INTEGER NOT NULL,
    event_type VARCHAR(32) NOT NULL,
    summary CLOB,
    created_at VARCHAR(64) NOT NULL,
    UNIQUE(session_id, turn_no)
);

CREATE TABLE agent_tool_call (
    id VARCHAR(64) PRIMARY KEY,
    session_id VARCHAR(64) NOT NULL,
    tool_call_id VARCHAR(128),
    tool_code VARCHAR(64) NOT NULL,
    arguments_json CLOB,
    result_summary CLOB,
    status VARCHAR(32) NOT NULL,
    duration_ms BIGINT NOT NULL,
    started_at VARCHAR(64) NOT NULL,
    finished_at VARCHAR(64) NOT NULL
);

CREATE TABLE agentscope_checkpoint (
    user_id VARCHAR(255) NOT NULL,
    session_id VARCHAR(255) NOT NULL,
    state_key VARCHAR(255) NOT NULL,
    item_index INTEGER NOT NULL,
    state_data CLOB NOT NULL,
    updated_at VARCHAR(64) NOT NULL,
    PRIMARY KEY (user_id, session_id, state_key, item_index)
);

COMMENT ON TABLE agent_session IS 'AgentScope 智能体运行会话';
COMMENT ON TABLE agent_turn IS '不包含隐藏思维过程的 AgentScope 可审计事件';
COMMENT ON TABLE agent_tool_call IS '智能体工具调用审计记录';
COMMENT ON TABLE agentscope_checkpoint IS 'AgentScope 可恢复运行状态';

INSERT INTO managed_agent (id, agent_code, agent_name, agent_type, description, endpoint_url, model_name, enabled,
                           runtime_status, last_heartbeat, total_tasks, success_tasks, failed_tasks, created_at, updated_at)
SELECT '00000000-0000-0000-0000-000000000405', 'OPS_REACT', '运维协调 Agent', 'REACT',
       '根据运维任务自动调度日志诊断、根因分析、本地 CLI 修复与验证，并在需要时请求人工审批', NULL, 'qwen-plus', 1, 'IDLE', NULL, 0, 0, 0,
       '2026-07-22T00:00:00Z', '2026-07-22T00:00:00Z'
FROM dual WHERE NOT EXISTS (SELECT 1 FROM managed_agent WHERE agent_code='OPS_REACT');

UPDATE managed_tool SET enabled=1, updated_at='2026-07-22T00:00:00Z' WHERE tool_code='LOG_QUERY';

INSERT INTO agent_tool_permission (id, agent_id, tool_id, timeout_seconds, daily_limit, enabled, created_at, updated_at)
SELECT '00000000-0000-0000-0000-000000000601', a.id, t.id, 60, 1000, 1,
       '2026-07-22T00:00:00Z', '2026-07-22T00:00:00Z'
FROM managed_agent a, managed_tool t
WHERE a.agent_code='OPS_REACT' AND t.tool_code='LOG_QUERY'
  AND NOT EXISTS (SELECT 1 FROM agent_tool_permission p WHERE p.agent_id=a.id AND p.tool_id=t.id);
'''.strip(),
        '',
        '-- ---------- D. 审批动作表（已合并指纹门禁列）----------',
        '''
CREATE TABLE agent_pending_action (
    id VARCHAR(64) PRIMARY KEY,
    session_id VARCHAR(64) NOT NULL,
    incident_id VARCHAR(64) NOT NULL,
    tool_call_id VARCHAR(128) NOT NULL,
    tool_code VARCHAR(64) NOT NULL,
    arguments_json CLOB,
    risk_level VARCHAR(16) NOT NULL,
    description VARCHAR(1000),
    status VARCHAR(16) NOT NULL,
    decision_comment VARCHAR(1000),
    requested_at VARCHAR(64) NOT NULL,
    decided_at VARCHAR(64),
    decided_by VARCHAR(64),
    patch_fingerprint VARCHAR(128),
    deployment_key VARCHAR(256),
    expires_at VARCHAR(64),
    invalidated_at VARCHAR(64)
);

CREATE INDEX idx_agent_pending_action_deployment
    ON agent_pending_action(incident_id, deployment_key, status);

COMMENT ON TABLE agent_pending_action IS 'ReAct Agent 等待人工决策的工具动作';

INSERT INTO managed_tool (id,tool_code,tool_name,tool_type,description,command_template,risk_level,timeout_seconds,enabled,created_at,updated_at)
SELECT '00000000-0000-0000-0000-000000000505','PATCH_PREVIEW','补丁预览','GIT','生成并检查补丁预览，不执行部署',NULL,'MEDIUM',120,1,'2026-07-22T00:00:00Z','2026-07-22T00:00:00Z'
FROM dual WHERE NOT EXISTS(SELECT 1 FROM managed_tool WHERE tool_code='PATCH_PREVIEW');

INSERT INTO managed_tool (id,tool_code,tool_name,tool_type,description,command_template,risk_level,timeout_seconds,enabled,created_at,updated_at)
SELECT '00000000-0000-0000-0000-000000000506','TEST_VERIFICATION','测试验证','TEST','执行受控测试并返回验证结果',NULL,'MEDIUM',600,1,'2026-07-22T00:00:00Z','2026-07-22T00:00:00Z'
FROM dual WHERE NOT EXISTS(SELECT 1 FROM managed_tool WHERE tool_code='TEST_VERIFICATION');

INSERT INTO managed_tool (id,tool_code,tool_name,tool_type,description,command_template,risk_level,timeout_seconds,enabled,created_at,updated_at)
SELECT '00000000-0000-0000-0000-000000000507','DEPLOY_PATCH','补丁部署','DEPLOY','部署已验证补丁，默认禁用并需单独授权',NULL,'CRITICAL',900,0,'2026-07-22T00:00:00Z','2026-07-22T00:00:00Z'
FROM dual WHERE NOT EXISTS(SELECT 1 FROM managed_tool WHERE tool_code='DEPLOY_PATCH');

INSERT INTO agent_tool_permission (id,agent_id,tool_id,timeout_seconds,daily_limit,enabled,created_at,updated_at)
SELECT '00000000-0000-0000-0000-000000000602',a.id,t.id,120,100,1,'2026-07-22T00:00:00Z','2026-07-22T00:00:00Z'
FROM managed_agent a, managed_tool t
WHERE a.agent_code='OPS_REACT' AND t.tool_code='PATCH_PREVIEW'
  AND NOT EXISTS(SELECT 1 FROM agent_tool_permission p WHERE p.agent_id=a.id AND p.tool_id=t.id);

INSERT INTO agent_tool_permission (id,agent_id,tool_id,timeout_seconds,daily_limit,enabled,created_at,updated_at)
SELECT '00000000-0000-0000-0000-000000000603',a.id,t.id,600,50,1,'2026-07-22T00:00:00Z','2026-07-22T00:00:00Z'
FROM managed_agent a, managed_tool t
WHERE a.agent_code='OPS_REACT' AND t.tool_code='TEST_VERIFICATION'
  AND NOT EXISTS(SELECT 1 FROM agent_tool_permission p WHERE p.agent_id=a.id AND p.tool_id=t.id);

INSERT INTO agent_approval_policy (id,agent_id,auto_approve_max_risk,production_requires_approval,high_risk_requires_approval,credential_access_enabled,created_at,updated_at)
SELECT '00000000-0000-0000-0000-000000000701',a.id,'LOW',1,1,0,'2026-07-22T00:00:00Z','2026-07-22T00:00:00Z'
FROM managed_agent a
WHERE a.agent_code='OPS_REACT'
  AND NOT EXISTS(SELECT 1 FROM agent_approval_policy p WHERE p.agent_id=a.id);
'''.strip(),
        '',
        '-- ---------- E. 本地修复工具 ----------',
        read_sql('codex-repair-migration.sql').rstrip(),
        '',
        '-- ---------- F. 知识库 ----------',
        read_sql('knowledge-base-migration.sql').rstrip(),
        '',
        'COMMIT;',
        '',
        '-- 初始化完成。默认管理员见平台 README（首次登录后请立即改密）。',
    ]
    return '\n\n'.join(parts) + '\n'


def write_upgrade_catalog() -> None:
    lines = [
        '智维超级员工 · 达梦存量库升级执行清单',
        '=' * 60,
        '说明：',
        '1. 仅用于「已有库」升级；新空库请执行 01-install-full.sql',
        '2. 增量脚本大多非幂等，执行前备份，已执行过的序号请跳过',
        '3. 源文件位于: super-employee-platform/super-employee-app/docs/sql/',
        '4. 建议在变更窗口由 DBA 按序执行，并填写执行记录',
        '',
        f'{"序号":<6}{"脚本文件":<55}{"说明"}',
        '-' * 100,
    ]
    for seq, name, desc in UPGRADE_ORDER:
        exists = 'OK' if (SQL_SRC / name).exists() else 'MISSING'
        lines.append(f'{seq:<6}{name:<55}{desc}  [{exists}]')
    lines.append('')
    lines.append('可选（测试环境）：')
    for seq, name, desc in OPTIONAL_SEED:
        lines.append(f'{seq:<6}{name:<55}{desc}')
    lines.append('')
    lines.append('执行记录（手工填写）：')
    lines.append(f'{"序号":<6}{"执行人":<12}{"日期":<14}{"结果":<10}{"备注"}')
    for seq, name, _ in UPGRADE_ORDER:
        lines.append(f'{seq:<6}{"":<12}{"":<14}{"":<10}')
    (DELIVERY_SQL / '02-upgrade-catalog.txt').write_text('\n'.join(lines) + '\n', encoding='utf-8')

    # 复制升级脚本到 delivery/sql/dameng/upgrade 便于打包带走
    up_dir = DELIVERY_SQL / 'upgrade'
    up_dir.mkdir(parents=True, exist_ok=True)
    for seq, name, desc in UPGRADE_ORDER + OPTIONAL_SEED:
        src = SQL_SRC / name
        if not src.exists():
            continue
        dest = up_dir / f'{seq}-{name}'
        shutil.copy2(src, dest)


def sync_source_schema_repair_engine() -> None:
    """Keep source schema-dameng in sync for future installs."""
    path = SQL_SRC / 'schema-dameng.sql'
    text = path.read_text(encoding='utf-8')
    if 'repair_engine' in text:
        return
    updated = text.replace(
        '    log_path VARCHAR(1024),\n',
        '    repair_engine VARCHAR(32),\n    log_path VARCHAR(1024),\n',
        1,
    )
    if updated != text:
        path.write_text(updated, encoding='utf-8')


# ---------- Word helpers ----------
def set_run_font(run, size=11, bold=False, color=SLATE, name='微软雅黑'):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)


def add_heading_cn(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size={1: 18, 2: 14, 3: 12}.get(level, 11), bold=True,
                     color=TEAL_DARK if level == 1 else TEAL)


def add_para(doc, text, size=11, bold=False, color=SLATE):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.clear()
        run = p.add_run(item)
        set_run_font(run, size=11)


def shade_header_row(row):
    for cell in row.cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), '0B7285')
        shd.set(qn('w:val'), 'clear')
        tc_pr.append(shd)
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = WHITE
                run.font.bold = True


def set_cell(cell, text, bold=False, size=10, color=SLATE):
    cell.text = ''
    run = cell.paragraphs[0].add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, bold=True, color=WHITE)
    shade_header_row(table.rows[0])
    for r_i, row in enumerate(rows, start=1):
        for c_i, val in enumerate(row):
            set_cell(table.rows[r_i].cells[c_i], val)
    doc.add_paragraph()


def new_doc(title: str, subtitle: str) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(title)
    set_run_font(r, size=22, bold=True, color=TEAL_DARK)
    s = doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run(subtitle)
    set_run_font(r, size=12, color=GRAY)
    return doc


def save_doc(doc: Document, name: str) -> Path:
    path = DELIVERY / name
    try:
        doc.save(str(path))
    except PermissionError:
        path = DELIVERY / name.replace('.docx', '-v2.docx')
        doc.save(str(path))
    return path


def build_delivery_list_doc() -> Path:
    doc = new_doc('智维超级员工', '最终成果与交付物清单 V1.2')
    add_heading_cn(doc, '一、本期最终成果摘要', 1)
    add_para(doc, '在既有「日志巡检 → 多 Agent 分析 → CLI 修码 → 验证 → 审批 → 部署」闭环基础上，完成本期能力收敛与交付固化：')
    add_bullets(doc, [
        '多技术栈工程画像：后端支持 Maven JAR / Java WAR(Tomcat) / Gradle / Node / Go / Python / Docker；端侧支持无端侧、Web(npm)、uni-app、微信原生小程序。',
        '混合自动部署：后端按栈路径替换或 WAR/Docker 专用部署器；Web 静态资源可同步；小程序仅验证不自动发布。',
        '修复引擎三选一 + auto：CodeX / Cursor / Claude；Agent 管理页并排环境检测与登录。',
        '知识库增强：支持 txt/md/docx 等文档导入；可根据文档或描述调用 Agent 预生成草稿，人工确认后入库。',
        '运维体验：列表默认每页 20 条、统计区压缩、表单字段说明、长弹窗内容区内部滚动。',
        '达梦交付基线：01-install-full 含工程画像与多栈字段；存量升级含 15b/15c/15d。',
    ])
    add_heading_cn(doc, '二、交付包结构', 1)
    add_para(doc, '本交付包位于项目根目录 delivery/，建议整体打包移交。路演/手册可在项目根目录或随包附送。')
    add_table(doc, ['路径', '内容', '对象'], [
        ('delivery/sql/dameng/01-install-full.sql', '达梦新环境一键初始化（整合版，含多栈字段）', 'DBA'),
        ('delivery/sql/dameng/02-upgrade-catalog.txt', '存量库升级顺序与执行记录表（含 15b–15d）', 'DBA'),
        ('delivery/sql/dameng/upgrade/*', '按序编号的增量 SQL 副本', 'DBA'),
        ('delivery/sql/other-db/*', 'MySQL/Oracle/Kingbase 全量 schema 副本（对齐需评估）', 'DBA/研发'),
        ('delivery/交付物清单.docx', '本文件：最终成果与交付索引', '项目/移交'),
        ('delivery/配置与部署清单.docx', '私有配置/Profile/依赖与多栈构建说明', '实施/运维'),
        ('delivery/上线验收清单.docx', '功能与安全验收条目（含多栈、知识库导入与 UI）', '项目/运维主管'),
        ('智维超级员工-产品路演.pptx / -v2.pptx', '路演 PPT', '售前/汇报'),
        ('智维超级员工-运维与管理人员操作手册.docx', '日常操作手册（含多栈、知识库导入生成）', '运维/管理员'),
        ('智维超级员工-路演操作手册.docx', '路演现场剧本', '售前/演示'),
        ('super-employee-platform/', '后端源码与可构建工程', '研发/实施'),
        ('super-employee-web/', '前端源码', '研发/实施'),
        ('测试工程/', '演示用故障应用（可选交付）', '售前/培训'),
    ])
    add_heading_cn(doc, '三、数据库交付说明', 1)
    add_bullets(doc, [
        '主基线：达梦（dameng）。新空库只执行 01-install-full.sql 一次。',
        '存量达梦库：按 02-upgrade-catalog.txt 顺序执行未跑过的脚本；多栈相关务必包含 15b→15c→15d。',
        '知识库导入/AI 预生成为应用层能力（POST /api/knowledge/generate），无额外 DDL。',
        'MySQL / Oracle / Kingbase：docs/sql 下有对应 schema-*.sql，近期运维字段增量以达梦为主，跨库交付需单独评估对齐。',
        '增量脚本非幂等：执行前备份；禁止对同一库重复执行已成功的 DDL。',
    ])
    add_heading_cn(doc, '四、建议移交检查', 1)
    add_bullets(doc, [
        '交付包文件齐全，SQL 可打开且编码为 UTF-8；upgrade 目录含 15d。',
        '提供默认管理员改密要求与生产密钥轮换要求。',
        '附上构建产物或构建命令（后端 jar / 前端 dist）。',
        '确认至少一种修复引擎 CLI 已安装并可登录；Agent 管理页环境检测为可用。',
        '确认 DashScope/模型配置可用（知识库 AI 预生成依赖）；不可用时可本地整理草稿。',
        '明确演示工程与多栈样例是否纳入本次交付范围。',
    ])
    return save_doc(doc, '交付物清单.docx')


def build_config_doc() -> Path:
    doc = new_doc('智维超级员工', '配置与部署清单 V1.2')
    add_heading_cn(doc, '一、部署组件', 1)
    add_table(doc, ['组件', '说明', '备注'], [
        ('数据库', '达梦 8（推荐）/ MySQL / Oracle / Kingbase', '先执行 SQL 再启动应用'),
        ('后端', 'Spring Boot（super-employee-app）', 'Profile：dev/prod 默认组含 dameng'),
        ('前端', 'Vue3（super-employee-web）', '构建后由 Nginx 或静态服务托管'),
        ('修复引擎', 'Cursor / CodeX / Claude CLI（至少一种）', '本机或构建机可执行；Agent 管理页可检测'),
        ('大模型网关', 'DashScope（AgentScope Model）', '故障分析与知识库 AI 预生成共用'),
        ('构建工具', '按任务工程画像：JDK17+/Maven/Gradle/Node/Go/Python/Docker', 'Java 制品构建勿用 JDK 8'),
    ])
    add_heading_cn(doc, '一附、工程画像与混合部署（任务级）', 1)
    add_table(doc, ['维度', '可选值', '说明'], [
        ('backendKind', 'MAVEN / JAVA_WAR / GRADLE / NODE / GO / PYTHON / DOCKER', '决定默认验证命令、制品发现与部署器'),
        ('clientKind', 'NONE / WEB_NPM / UNI_APP / MP_WEIXIN_NATIVE', 'Web 可配静态部署目录；小程序仅验证'),
        ('部署方式', '路径替换 / WAR→webapps / Docker Compose / 静态同步', '由 PatchDeploymentFacade 按画像路由'),
        ('关键字段', 'backendVerifyCommand、clientDeployPath、clientDistDir、dockerComposeFile 等', '见升级脚本 15c/15d'),
    ])
    add_heading_cn(doc, '二、数据库初始化', 1)
    add_bullets(doc, [
        '新环境：执行 delivery/sql/dameng/01-install-full.sql',
        '存量升级：按 delivery/sql/dameng/02-upgrade-catalog.txt',
        '应用侧 spring.sql.init.mode 一般为 never，需手工执行 SQL',
        '开启持久化：super-employee.persistence.enabled=true',
    ])
    add_heading_cn(doc, '三、关键配置项', 1)
    add_table(doc, ['配置/私有文件', '含义', '建议'], [
        ('application-private.yml', '数据库连接凭据与安全密钥', '本地私有配置，已忽略，勿入库'),
        ('REACT_ENABLED / AI_ENABLED', '多 Agent 流程开关', '生产按就绪度开启'),
        ('SSH_DIAGNOSTIC_ENABLED', 'SSH 只读诊断', '生产显式开启并配指纹'),
        ('SSH_HOST_KEY_FINGERPRINTS', '主机指纹映射', '或在主机管理页面配置'),
        ('OPS_LOG_SCHEDULE_ENABLED', '运维任务巡检调度', '默认开启'),
        ('CODEX_REPAIR_ENABLED 等', '各 CLI 修复引擎开关', '按实际安装开启'),
        ('LOCAL_REPAIR_ENGINE', '默认引擎 codex/cursor/claude', '可与任务级配置配合'),
        ('CODEX/CURSOR/CLAUDE_WORKSPACE', '默认工作区', '任务也可单独指定'),
        ('ARTIFACT_JAVA_HOME 等', '制品构建 JDK', '指向 JDK 17+'),
        ('WEB_ORIGIN', '前端来源 CORS', '配置真实前端域名'),
        ('CAPTCHA_* / LOGIN_*', '验证码与登录锁定', '按安全基线设置'),
        ('agentscope.dashscope.api-key', 'DashScope API Key', '知识库预生成与模型分析依赖'),
        ('agentscope.dashscope.model-name', '模型名，默认 qwen-plus', '可按配额调整'),
    ])
    add_heading_cn(doc, '三附、知识库导入与预生成', 1)
    add_bullets(doc, [
        '接口：POST /api/knowledge/generate（multipart：file 与/或 description）',
        '支持：txt/md/docx/log/json/yaml/html 等；PDF/旧版 Office 请先另存为文本或 docx',
        '流程：预生成草稿 → 前端填入表单 → 人工确认后走原有保存接口入库',
        '模型不可用时回退本地模板整理，不阻断录入',
    ])
    add_heading_cn(doc, '四、推荐启动顺序', 1)
    add_bullets(doc, [
        '1）准备数据库并执行整合 SQL',
        '2）填写本地 application-private.yml 中的数据库与密钥配置',
        '3）启动后端并检查健康状态',
        '4）部署/启动前端并登录',
        '5）录入凭据与主机 → 创建运维任务 → 小流量试跑',
    ])
    add_heading_cn(doc, '五、安全基线', 1)
    add_bullets(doc, [
        '首次登录立即修改管理员密码',
        '生产关闭不必要的 demo 种子数据',
        'DEPLOY_PATCH 等高危工具默认禁用，按需授权',
        '日志白名单最小化；禁止放开任意远程 Shell',
    ])
    return save_doc(doc, '配置与部署清单.docx')


def build_acceptance_doc() -> Path:
    doc = new_doc('智维超级员工', '上线验收清单 V1.2')
    add_heading_cn(doc, '一、基础可用性', 1)
    add_table(doc, ['编号', '验收项', '通过标准', '结果'], [
        ('A01', '登录', '验证码登录成功，可改密', ''),
        ('A02', '菜单权限', '管理员可见核心菜单；无权限账号被正确拦截', ''),
        ('A03', '健康检查', '后端健康检查正常', ''),
        ('A04', '审计', '关键操作可在审计/执行记录中追溯', ''),
        ('A05', '列表与表单体验', '列表默认可展示更多行；长表单弹窗内部滚动；关键字段有说明', ''),
    ])
    add_heading_cn(doc, '二、资源与任务', 1)
    add_table(doc, ['编号', '验收项', '通过标准', '结果'], [
        ('B01', '凭据管理', '新增 SSH 凭据，密钥不回显', ''),
        ('B02', '主机纳管', '指纹+日志白名单配置成功，探测连通性通过并启用', ''),
        ('B03', '运维任务（多栈）', '可选择 backendKind/clientKind，工作区与验证命令保存成功', ''),
        ('B04', '巡检拉日志', '启用后能拉取日志并生成摘要/落盘', ''),
        ('B05', '修复引擎环境', 'Agent 管理页可检测 CodeX/Cursor/Claude 安装与登录态', ''),
    ])
    add_heading_cn(doc, '三、自治闭环', 1)
    add_table(doc, ['编号', '验收项', '通过标准', '结果'], [
        ('C01', '异常建故障', '错误达到阈值后派生故障事件', ''),
        ('C02', '多 Agent 分析', '执行记录可见证据/根因/规划阶段', ''),
        ('C03', '补丁预览', '生成可审阅补丁计划与指纹', ''),
        ('C04', '修复审批', '未审批不能修码；审批后 CLI 产出真实 Diff', ''),
        ('C05', '验证构建', '按工程画像完成验证（Maven/Gradle/npm 等）并得到对应制品', ''),
        ('C06', '部署门禁', '部署前需再次审批/确认', ''),
        ('C07', '混合部署（可选）', '后端路径替换或 WAR/Docker 部署成功；Web 静态同步按配置生效', ''),
    ])
    add_heading_cn(doc, '四、治理与安全', 1)
    add_table(doc, ['编号', '验收项', '通过标准', '结果'], [
        ('D01', '只读诊断', '远程仅固定模板，不能任意 Shell', ''),
        ('D02', '授权策略', '工具/资源授权生效', ''),
        ('D03', '知识库 CRUD', '可新增/编辑/查询；不可用时不阻断主流程', ''),
        ('D04', '知识库导入预生成', '上传 txt/md/docx 或填写描述可生成草稿；确认后保存入库', ''),
        ('D05', '驳回路径', '审批驳回后流程停止或回到可调整状态', ''),
    ])
    add_heading_cn(doc, '五、签核', 1)
    add_table(doc, ['角色', '姓名', '日期', '签字'], [
        ('实施负责人', '', '', ''),
        ('运维负责人', '', '', ''),
        ('甲方项目负责人', '', '', ''),
    ])
    return save_doc(doc, '上线验收清单.docx')


def write_readme() -> None:
    text = '''# 智维超级员工 · 交付包（delivery）

版本：V1.2（多技术栈混合部署 + 知识库导入预生成 + 运维体验收敛）

## 本期成果要点

- 后端工程画像：MAVEN / JAVA_WAR / GRADLE / NODE / GO / PYTHON / DOCKER
- 端侧画像：NONE / WEB_NPM / UNI_APP / MP_WEIXIN_NATIVE（Web 可静态同步，小程序仅验证）
- 存量升级关键脚本：`15b` 前端目录 → `15c` 工程画像 → `15d` 多栈部署字段
- 知识库：文档导入（txt/md/docx 等）+ Agent 预生成草稿，人工确认后入库
- 体验：列表默认 20 条、表单字段说明、长弹窗内部滚动、Agent 三引擎环境检测

## 目录说明

| 文件/目录 | 用途 |
|-----------|------|
| sql/dameng/01-install-full.sql | 达梦**新空库**一键初始化（整合版，仅执行一次） |
| sql/dameng/02-upgrade-catalog.txt | 达梦**存量库**升级顺序与执行记录 |
| sql/dameng/upgrade/ | 按序号复制的增量 SQL，便于离线移交 |
| sql/other-db/ | 其他库全量 schema 副本（对齐需评估） |
| 交付物清单.docx | 最终成果与全部可交付索引 |
| 配置与部署清单.docx | 部署组件与关键配置项 |
| 上线验收清单.docx | 上线/移交验收表 |

## 数据库怎么用

### 新环境（推荐）
1. 创建达梦用户/模式
2. 执行 `sql/dameng/01-install-full.sql`
3. 在本地 `local-config/application-private.yml` 配置数据库与密钥后启动

### 已有环境升级
1. 备份数据库
2. 打开 `02-upgrade-catalog.txt`，跳过已执行脚本
3. 按序执行 `upgrade/` 下未执行文件（多栈务必含 15b→15c→15d）
4. 填写执行记录

### 其他数据库
MySQL / Oracle / Kingbase 全量脚本见 `sql/other-db/` 与源码 `docs/sql/schema-*.sql`。
近期增量以达梦为主，跨库交付请单独评估。

## 重新生成

```powershell
python .\\scripts\\build_delivery_package.py
python .\\scripts\\generate_ops_manual.py
python .\\scripts\\generate_pitch_deck.py
python .\\scripts\\generate_pitch_ops_doc.py
```

项目根目录同步生成/更新：
- `智维超级员工-产品路演.pptx` / `-v2.pptx`
- `智维超级员工-运维与管理人员操作手册.docx`
- `智维超级员工-路演操作手册.docx`
'''
    (DELIVERY / 'README.md').write_text(text, encoding='utf-8')


def main() -> None:
    DELIVERY_SQL.mkdir(parents=True, exist_ok=True)
    sync_source_schema_repair_engine()
    install = build_install_full()
    (DELIVERY_SQL / '01-install-full.sql').write_text(install, encoding='utf-8')
    write_upgrade_catalog()
    write_readme()

    # 其它库 schema 只读副本索引
    other = DELIVERY / 'sql' / 'other-db'
    other.mkdir(parents=True, exist_ok=True)
    note = [
        '以下为 docs/sql 中的全量脚本副本，便于打包。',
        '注意：近期运维任务/知识库/Agent 运行时增量以达梦整合包为准；',
        'MySQL/Oracle/Kingbase 需额外评估对齐后再用于生产。',
        '',
    ]
    for name in ['schema-mysql.sql', 'schema-oracle.sql', 'schema-kingbase.sql', 'schema.sql']:
        src = SQL_SRC / name
        if src.exists():
            shutil.copy2(src, other / name)
            note.append(f'- {name}')
    (other / 'README.txt').write_text('\n'.join(note) + '\n', encoding='utf-8')

    p1 = build_delivery_list_doc()
    p2 = build_config_doc()
    p3 = build_acceptance_doc()

    print('DELIVERY=', DELIVERY)
    print('SQL_FULL=', DELIVERY_SQL / '01-install-full.sql')
    print('DOCS=', p1.name, p2.name, p3.name)
    print('upgrade_scripts=', len(list((DELIVERY_SQL / 'upgrade').glob('*.sql'))))


if __name__ == '__main__':
    main()
