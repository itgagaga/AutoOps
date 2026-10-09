# -*- coding: utf-8 -*-
"""生成《智维超级员工-AI提效说明》Word，供内部汇报与试点对照。"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TEAL = RGBColor(0x0B, 0x72, 0x85)
TEAL_DARK = RGBColor(0x0B, 0x4F, 0x5C)
SLATE = RGBColor(0x1F, 0x2A, 0x37)
GRAY = RGBColor(0x52, 0x60, 0x6D)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run_font(run, size=11, bold=False, color=SLATE, name='微软雅黑'):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)


def add_heading_cn(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(
            run,
            size={1: 18, 2: 14, 3: 12}.get(level, 11),
            bold=True,
            color=TEAL_DARK if level == 1 else TEAL,
        )
    return p


def add_para(doc, text, size=11, bold=False, color=SLATE, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_bullets(doc, items, size=11):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.clear()
        run = p.add_run(item)
        set_run_font(run, size=size, color=SLATE)
        p.paragraph_format.space_after = Pt(3)


def add_numbered(doc, items, size=11):
    for item in items:
        p = doc.add_paragraph(style='List Number')
        p.clear()
        run = p.add_run(item)
        set_run_font(run, size=size, color=SLATE)
        p.paragraph_format.space_after = Pt(4)


def set_cell_text(cell, text, bold=False, color=SLATE, size=10):
    cell.text = ''
    p = cell.paragraphs[0]
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)


def shade_header_row(row):
    for cell in row.cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:fill'), '0B7285')
        shd.set(qn('w:val'), 'clear')
        tc_pr.append(shd)
        for p in cell.paragraphs:
            for run in p.runs:
                run.font.color.rgb = WHITE
                run.font.bold = True


def add_table(doc, headers, rows):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.style = 'Table Grid'
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, size=10)
    shade_header_row(table.rows[0])
    for r_i, row in enumerate(rows, start=1):
        for c_i, val in enumerate(row):
            set_cell_text(table.rows[r_i].cells[c_i], val, size=10)
    doc.add_paragraph()
    return table


def main():
    root = Path(__file__).resolve().parents[1]
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(36)
    r = t.add_run('智维超级员工')
    set_run_font(r, size=28, bold=True, color=TEAL_DARK)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run('AI 提效说明')
    set_run_font(r, size=20, bold=True, color=TEAL)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run('用多 Agent 把故障处置从「人肉串岗」变成「人在回路的受控提速」')
    set_run_font(r, size=12, color=GRAY)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(16)
    r = meta.add_run('版本：V1.2 ｜ 适用：管理层、运维主管、试点评估 ｜ 内部资料')
    set_run_font(r, size=10, color=GRAY)

    add_para(
        doc,
        '本文说明平台如何把 AI 用在真实运维闭环上：缩短发现与定位时间、压缩修码与验证等待、'
        '把经验沉淀为可复用知识，同时用审批与审计把风险锁住。文中工时为演练工程与同类试点的对照口径，'
        '落地后请用本单位工单数据替换。',
        size=10,
        color=GRAY,
    )

    add_heading_cn(doc, '一、提效结论（先看这一页）', 1)
    add_para(
        doc,
        '传统处置依赖「值班发现 → 人肉翻日志 → 找研发 → 改代码 → 等测试 → 找窗口发布」。'
        '智维超级员工把可标准化的前半段交给多 Agent，人只做高风险确认。',
    )
    add_table(doc, ['指标', '传统口径（单次常见故障）', '平台介入后', '提效方向'], [
        ('发现时延', '依赖告警或用户投诉，数十分钟到数小时', '按任务周期拉日志，分钟级发现', '早发现'),
        ('定位根因', '跨人、跨系统，约 30–90 分钟', '证据收集 + 根因 Agent，约 5–15 分钟', '少串岗'),
        ('给出补丁', '研发排期，数小时到隔日', '本地修复引擎生成 diff，约 5–20 分钟', '快出方案'),
        ('验证等待', '人工搭环境、跑用例，约 20–60 分钟', '按工程画像自动验证', '少等待'),
        ('发布风险', '口口相传、易漏检', '补丁指纹绑定 + 二次审批', '控得住'),
        ('经验复用', '个人笔记，难传承', '知识库导入 + AI 预生成草稿', '可沉淀'),
    ])
    add_para(
        doc,
        '综合口径：同类「日志可定位、代码可修」的运行时缺陷，单次闭环人工耗时可由约 2–4 小时压到约 20–40 分钟'
        '（含两次人工审批）。值班人力从「盯日志」转为「审 diff、把窗口」。',
        bold=True,
    )

    add_heading_cn(doc, '二、痛点：时间都耗在哪里', 1)
    add_para(doc, '运维效率低，通常不是「不会修」，而是环节断裂：')
    add_bullets(doc, [
        '发现晚：没有持续拉日志，或日志散落多机，等用户报障才开工。',
        '定位慢：人在 SSH、日志文件、知识文档、代码仓库之间来回切换。',
        '协作贵：运维写现象、研发猜根因、测试等包，一次故障多次转派。',
        '修码不可控：有人直接改生产、或修完无测试、或补丁对不上现场版本。',
        '经验流失：同类 NPE / 解析错误反复出现，每次从零排查。',
    ])
    add_para(
        doc,
        'AI 如果只做聊天问答，解决不了上述问题。必须把「采集、分析、修码、验证、部署」串成一条带门禁的流水线。',
    )

    add_heading_cn(doc, '三、AI 用在哪几段（不是替代运维）', 1)
    add_table(doc, ['环节', '谁做', 'AI / Agent 做什么', '人必须做什么'], [
        ('日志巡检', '平台调度', '周期拉取、异常判定、派生故障事件', '配置主机、白名单、周期'),
        ('证据收集', '日志提取 Agent', '按模板取 HOST_HEALTH / JOURNAL / APP_LOG', '授权主机与工具'),
        ('根因分析', '根因 Agent', '结合日志与知识库给出因果与影响面', '确认是否属实'),
        ('修复规划', '协调 Agent', '生成补丁预览与文件白名单', '审阅范围'),
        ('本地修码', 'CodeX / Cursor / Claude', '在工作区写最小 diff', '审批启动 CLI'),
        ('测试验证', '验证 Agent + 构建', '按 Maven/多栈命令跑验证', '看失败则驳回'),
        ('部署发布', '状态机 + 部署器', '按画像替换制品 / WAR / 静态资源', '变更窗口内二次审批'),
        ('知识沉淀', '知识库 Agent', '文档导入后预生成条目草稿', '确认后入库'),
    ])
    add_para(doc, '原则：证据与执行留在内网；模型只看脱敏上下文；有副作用的修码和部署必须人点头。')

    add_heading_cn(doc, '四、分环节对照（演练口径）', 1)
    add_heading_cn(doc, '4.1 发现与派单', 2)
    add_bullets(doc, [
        '传统：靠监控阈值或群里报障；漏报夜间故障。',
        '平台：运维任务按 5/10/30/60 分钟拉日志，异常达到阈值自动建事件。',
        '提效：值班从「主动翻机」变为「看待办」；夜间可先出分析草稿，早班审批。',
    ])
    add_heading_cn(doc, '4.2 定位与方案', 2)
    add_bullets(doc, [
        '传统：打开 journal、应用日志、配置，口头同步研发。',
        '平台：固定只读 SSH 模板 + 多 Agent 时间线，根因与补丁计划写进事件。',
        '提效：减少「把日志打包发给研发」的往返；定位从小时级到分钟级（可复现缺陷）。',
    ])
    add_heading_cn(doc, '4.3 修码与验证', 2)
    add_bullets(doc, [
        '传统：等研发排期；本地改完再找测试环境。',
        '平台：审批后 CLI 只改白名单文件，随后自动 Maven/多栈验证，diff 可回看。',
        '提效：常见空指针、格式解析、越界类缺陷可当场出补丁；验证不再单独占一个工位半天。',
    ])
    add_heading_cn(doc, '4.4 发布与复盘', 2)
    add_bullets(doc, [
        '传统：U 盘拷包、口头确认版本，复盘记在个人文档。',
        '平台：补丁指纹绑定验证结果；部署再审批；完结后回填知识库。',
        '提效：同类故障二次出现时可检索知识，缩短「从零分析」时间。',
    ])

    add_heading_cn(doc, '五、典型场景与人力变化', 1)
    add_table(doc, ['场景', '传统做法', '用平台之后', '对人的要求'], [
        ('应用持续打 ERROR', '登录主机 grep，群里@研发', '任务自动拉日志并分析', '配好路径与周期'),
        ('空指针 / 解析失败', '开 IDE 猜，改完再测', 'CLI 出 diff，自动跑测试', '审差异、点批准'),
        ('WAR / 前后端混合项目', '分别找人构建、分别发布', '按工程画像验证并混合部署', '绑对部署目录'),
        ('值班交接', '口头交代未完故障', '事件时间线 + 审计可续跑', '看状态机即可接班'),
        ('制度与经验宣贯', '开会讲、发邮件', '知识库导入制度文档，分析时可引用', '审核 AI 草稿'),
    ])
    add_para(doc, '人力结构建议：', bold=True)
    add_bullets(doc, [
        '一线值班：从「搜日志」转为「盯任务列表、处理审批、确认完结」。',
        '二线 / 研发：从「每次救火」转为「审高风险补丁、处理 Agent 搞不定的疑难」。',
        '主管：从「追进度」转为「看审批队列、变更窗口与审计」。',
    ])

    add_heading_cn(doc, '六、演示工程可复现的提效路径', 1)
    add_para(
        doc,
        '仓库内「测试工程」（demo-ops-app）内置可触发缺陷：NPE、NumberFormat、日期解析、越界、'
        '缺失模板文件、短 SKU、并发修改等。用于证明「从日志到补丁」可在单次路演内走完。',
    )
    add_numbered(doc, [
        '启动演示应用，确认 logs/demo-ops-app.log 持续出现 ERROR（或运行 scripts/trigger-faults.ps1）。',
        '平台运维任务指向该日志路径与代码工作区，开启自动分析。',
        '观察故障事件时间线：提取 → 根因 → 补丁预览。',
        '审批本地修复，审阅 diff，接受后自动验证。',
        '变更窗口内审批部署（演示可用受控命令代替真实生产发布）。',
        '将根因与处置写入知识库，作为下次同类故障的上下文。',
    ])
    add_para(
        doc,
        '路演对照：人工逐步触发并口头讲解约需 30–40 分钟；平台自动巡检 + 操作员只点审批，闭环可压到 8–15 分钟。',
    )

    add_heading_cn(doc, '七、怎么衡量（建议试点 4 周）', 1)
    add_para(doc, '不要只报「用了 AI」，按工单打点。建议至少统计：')
    add_table(doc, ['KPI', '统计方式', '试点目标（建议）'], [
        ('MTTD 平均发现时间', '异常首次出现 → 平台建事件', '较基线下降 ≥ 50%'),
        ('MTTI 平均定位时间', '建事件 → 根因结论可审阅', '可复现类下降 ≥ 60%'),
        ('MTTR 平均恢复时间', '建事件 → 验证通过或部署完成', '可自动修类别下降 ≥ 40%'),
        ('人工触点数', '一次故障需要的沟通/转派次数', '降至审批 1–2 次'),
        ('补丁一次通过率', '验证失败退回次数', '试点期记录基线后逐周改善'),
        ('知识命中率', '分析时引用知识库条目占比', '第 4 周可见上升'),
        ('越权/误发布次数', '审计中未走审批的变更', '目标为 0'),
    ])
    add_para(
        doc,
        '口径说明：仅统计「日志可定位且工作区可编译验证」的故障；基础设施宕机、无源代码、'
        '需变更窗口外紧急热修的，单独分类，不纳入自动修复成功率。',
        size=10,
        color=GRAY,
    )

    add_heading_cn(doc, '八、提效不等于失控：安全与治理', 1)
    add_bullets(doc, [
        'SSH 只有只读模板，模型不能下发任意 Shell。',
        '本地修复必须在补丁预览之后，且每次启动 CLI 需审批。',
        '验证绑定真实 diff / 补丁指纹，防止「分析的是 A、部署的是 B」。',
        '部署再次审批；高危工具默认关闭，按 Agent 授权放开。',
        '工具调用、会话、审批全程可审计，便于等保与内部稽核。',
        '生产关闭 dev 默认口令与开发 JWT；密钥走环境变量。',
    ])
    add_para(doc, '对管理层的一句话：速度来自自动化前半段；安全来自人在回路的后半段。')

    add_heading_cn(doc, '九、落地节奏（避免一上来全自动）', 1)
    add_table(doc, ['阶段', '周期', '做法', '提效预期'], [
        ('只读观察', '第 1 周', '只开巡检与分析，关闭自动修码', '先缩短发现与定位'),
        ('修码试点', '第 2–3 周', '对演示/预发应用开 CLI 修复 + 验证，部署仍人工', '压缩方案与验证时间'),
        ('受控发布', '第 4 周', '预发允许自动部署；生产保持二次审批', '形成完整闭环'),
        ('知识运营', '持续', '每周把完结事件回填知识库', '同类故障越来越快'),
    ])
    add_para(doc, '不建议第一步就对核心生产开启自动部署。先用预发和演示工程把指纹、工作区、CLI 跑稳。')

    add_heading_cn(doc, '十、常见误区', 1)
    add_table(doc, ['误区', '正确理解'], [
        ('上了大模型就等于自动运维', '没有日志白名单、工作区和审批，模型只能聊天，不能闭环'),
        ('AI 会替换运维编制', '替换的是重复翻日志和等待；编制转向治理与疑难'),
        ('一次修好所有故障', '只覆盖可复现、有代码工作区的应用缺陷；硬件/网络另案'),
        ('越快越好、审批是阻力', '两次审批换来可审计与可回滚，是提效可持续的前提'),
        ('知识库可自动全量入库', 'AI 只出草稿，人工确认后才成为分析上下文'),
    ])

    add_heading_cn(doc, '十一、与本文相关的材料', 1)
    add_table(doc, ['材料', '用途'], [
        ('README.md「生产部署」', '控制台怎么装、怎么开 prod'),
        ('智维超级员工-运维与管理人员操作手册.docx', '日常怎么配任务、怎么审批'),
        ('智维超级员工-路演操作手册.docx', '8–15 分钟现场走通闭环'),
        ('智维超级员工-成果汇报.docx', '能力清单与截图位'),
        ('delivery/上线验收清单.docx', '上线是否达到安全与功能基线'),
    ])

    add_heading_cn(doc, '十二、版本说明', 1)
    add_table(doc, ['项', '内容'], [
        ('产品版本', 'V1.2'),
        ('文档用途', 'AI 提效评估 / 试点汇报 / 管理沟通'),
        ('数据性质', '演练与同类试点对照口径，需用本单位工单校准'),
        ('编制', '与 delivery 脚本一同维护，可重复生成'),
    ])

    out = root / '智维超级员工-AI提效说明.docx'
    try:
        doc.save(str(out))
    except PermissionError:
        out = root / '智维超级员工-AI提效说明-v2.docx'
        doc.save(str(out))
    print(f'已生成: {out}')
    return out


if __name__ == '__main__':
    main()
